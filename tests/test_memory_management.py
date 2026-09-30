"""Regresiones de paginación, lotes, recursos y descargas grandes."""
from contextlib import ExitStack
from datetime import date
from io import BytesIO
from tempfile import TemporaryFile
from threading import Event, Thread
import unittest
from unittest.mock import patch

from flask import template_rendered
from openpyxl import Workbook
from sqlalchemy import create_engine, event, inspect
from sqlalchemy.orm import sessionmaker
import xlsxwriter

from tests.test_integration import app
from tests.csv_helpers import excel_csv_reader
from tests.test_order_import import order, workbook_bytes
import app as main
from gestion5s import web
from gestion5s.orders import parse_orders
from gestion5s.order_import_routes import _insert_groups
from memory_utils import DiskRows, memory_limited, send_disk_file


class MemoryManagementTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        web.Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.session_patch = patch.object(web, "SessionLocal", self.sessions)
        self.session_patch.start()
        self.client = app.test_client()

    def tearDown(self):
        self.session_patch.stop()
        self.engine.dispose()

    def test_all_record_categories_paginate_without_truncating_csv(self):
        for entity, values in (
            ("desviaciones", {"fecha": date(2026, 9, 1)}),
            ("entradas_salidas", {"fecha_ingreso": None}),
        ):
            model = web.ENTITY_MODEL[entity]
            with self.sessions() as db:
                db.execute(model.__table__.insert(), [dict(values) for _ in range(205)])
                db.commit()
            contexts = []
            def capture(sender, template, context, **extra):
                if template.name == "list.html":
                    contexts.append(context)
            with template_rendered.connected_to(capture, web.app):
                for page, expected in ((1, 100), (2, 100), (3, 5)):
                    with self.client.get(f"/gestion-5s/registros?vista={entity}&page={page}") as response:
                        self.assertEqual(response.status_code, 200)
                    self.assertEqual(len(contexts[-1]["current_records"]), expected)
                    self.assertEqual(contexts[-1]["record_count"], 205)
                ids = [row.id for ctx in contexts for row in ctx["current_records"]]
                self.assertEqual(len(set(ids)), 205)
            with self.client.get(f"/gestion-5s/download/{entity}.csv") as response:
                self.assertEqual(len(list(excel_csv_reader(response.data))), 205)

    def test_generic_import_flushes_batches_and_rolls_back_late_invalid_row(self):
        entity = "habitaciones_bloqueadas"
        model = web.ENTITY_MODEL[entity]
        with self.sessions() as db:
            db.add(model(habitacion="PREVIA"))
            db.commit()
        def upload(invalid=False):
            book = Workbook(write_only=True)
            sheet = book.create_sheet()
            sheet.append(web.TEMPLATES[entity])
            for i in range(1001):
                sheet.append(["2026-09-01", f"{i:05}"] + [None] * 10)
            if invalid:
                sheet.append(["fecha-inválida", "ERROR"] + [None] * 10)
            buffer = BytesIO()
            book.save(buffer)
            buffer.seek(0)
            return self.client.post(f"/gestion-5s/import/{entity}", data={"file": (buffer, "plantilla.xlsx")})
        pending = []
        def before_flush(session, context, instances):
            pending.append(len(session.new))
        event.listen(self.sessions, "before_flush", before_flush)
        try:
            self.assertEqual(upload(invalid=True).status_code, 302)
            with self.sessions() as db:
                self.assertEqual(db.query(model).count(), 1)
                self.assertEqual(db.query(model).one().habitacion, "PREVIA")
            self.assertEqual(upload().status_code, 302)
            with self.sessions() as db:
                self.assertEqual(db.query(model).count(), 1002)
            self.assertTrue(pending)
            self.assertLessEqual(max(pending), 500)
        finally:
            event.remove(self.sessions, "before_flush", before_flush)

    def test_disk_order_batches_keep_every_row_and_date_type(self):
        records = [order(f"{i:06}", "CARPINTERIA MENOR" if i % 2 else "GASFITER") for i in range(1203)]
        with ExitStack() as resources:
            stores = []
            def factory():
                store = resources.enter_context(DiskRows())
                stores.append(store)
                return store
            parsed = parse_orders(workbook_bytes(records), details_factory=factory)
            self.assertEqual(parsed["total"], 1203)
            self.assertEqual(parsed["start"], date(2026, 8, 31))
            self.assertEqual(parsed["groups"]["solicitud_ot"][:1][0]["ot"], "000000")
            with self.sessions() as db:
                _insert_groups(db, web.ENTITY_MODEL, parsed["groups"])
                db.commit()
                self.assertEqual(db.query(web.MiscelaneoEntry).count(), 601)
                self.assertEqual(db.query(web.SolicitudOTEntry).count(), 602)
            self.assertEqual(len(list(parsed["groups"]["solicitud_ot"])), 602)
            self.assertEqual(len(list(parsed["groups"]["solicitud_ot"])), 602)
        self.assertTrue(all(store.file.closed for store in stores))

    def test_download_closes_temporaries_after_success_or_disconnect(self):
        for consume in (True, False):
            output = TemporaryFile(mode="w+b")
            output.write(b"x" * 20000)
            output.seek(0)
            with app.test_request_context():
                response = send_disk_file(output, download_name="test.csv", mimetype="text/csv")
                if consume:
                    self.assertEqual(b"".join(response.response), b"x" * 20000)
                else:
                    response.response.close()
                self.assertTrue(output.closed)

    def test_second_heavy_operation_is_rejected_and_lock_recovers(self):
        entered, release = Event(), Event()
        @memory_limited()
        def first():
            entered.set()
            release.wait(timeout=10)
        worker = Thread(target=first)
        worker.start()
        try:
            self.assertTrue(entered.wait(timeout=3))
            with self.client.get("/reports/usuarios.xlsx?start_date=2026-09-01&end_date=2026-09-30") as response:
                self.assertEqual(response.status_code, 503)
                self.assertEqual(response.headers["Retry-After"], "30")
        finally:
            release.set()
            worker.join(timeout=3)
        @memory_limited()
        def failing():
            raise ValueError("fallo de prueba")
        with self.assertRaises(ValueError):
            failing()
        with self.client.get("/reports/usuarios.xlsx?start_date=2026-09-01&end_date=2026-09-30") as response:
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.data.startswith(b"PK"))

    def test_metadata_queries_defer_binary_payloads_but_download_can_load_them(self):
        with app.app_context():
            upload = main.UploadedFile(filename="memory-test.xlsx", file_type="curva", sha256="test", content=b"SOURCE")
            job = main.ExportJob(status="completed", content=b"EXPORT")
            main.db.session.add_all([upload, job])
            main.db.session.commit()
            ids = upload.id, job.id
            main.db.session.expunge_all()
            try:
                for model, key, expected in ((main.UploadedFile, ids[0], b"SOURCE"), (main.ExportJob, ids[1], b"EXPORT")):
                    record = main.db.session.get(model, key)
                    self.assertIn("content", inspect(record).unloaded)
                    self.assertEqual(record.content, expected)
            finally:
                main.ExportJob.query.filter_by(id=ids[1]).delete()
                main.UploadedFile.query.filter_by(id=ids[0]).delete()
                main.db.session.commit()

    def test_xml_reader_ignores_formatted_empty_rows_and_preserves_header_offsets(self):
        output = BytesIO()
        with xlsxwriter.Workbook(output) as book:
            sheet = book.add_worksheet("Censo")
            sheet.write(0, 0, "Título")
            sheet.write_blank(1, 0, None, book.add_format({"bold": True}))
            sheet.write_row(2, 0, ["Id", "Camas Ocupadas", "Comentario"])
            sheet.write_row(3, 0, ["000012", 0, "área A"])
            blank = book.add_format({"bg_color": "#EEEEEE"})
            for row in range(4, 5004):
                sheet.write_blank(row, 16383, None, blank)
        content = output.getvalue()
        raw = main.read_xlsx_sheet_df(content, "Censo", nrows=4)
        self.assertEqual(raw.iloc[2, 0], "Id")
        frame = main.read_xlsx_sheet_df(content, "Censo", header=2)
        self.assertEqual(frame.shape, (1, 3))
        self.assertEqual(frame.iloc[0].tolist(), ["000012", 0, "área A"])

    def test_censo_rebuild_counts_all_batches_and_rolls_back_on_failure(self):
        import pandas as pd
        with app.app_context():
            upload = main.UploadedFile(filename="censo-memory.xlsx", file_type="censo", sha256="test", content=b"SOURCE")
            main.db.session.add(upload)
            main.db.session.flush()
            censo = main.Censo(file_id=upload.id, fecha_censo=date(2026, 9, 1))
            main.db.session.add(censo)
            main.db.session.flush()
            ids = censo.id, upload.id
            frame = pd.DataFrame({"Id": [f"R{i:05}" for i in range(1003)],
                                  "Camas Ocupadas": [i % 2 for i in range(1003)]})
            try:
                with patch.object(main, "lookup_active_curve_items", return_value=(None, {})):
                    main.rebuild_censo_from_dataframe(censo, frame)
                    main.db.session.commit()
                    self.assertEqual((censo.total_records, censo.total_occupied, censo.unmatched_count), (1003, 501, 1003))
                    self.assertEqual(main.CensoRecord.query.filter_by(censo_id=censo.id).count(), 1003)
                    original = main.norm_id
                    def invalid_after_first_batch(value):
                        if value == "R00900":
                            raise ValueError("error después de guardar el primer lote")
                        return original(value)
                    with patch.object(main, "norm_id", side_effect=invalid_after_first_batch):
                        with self.assertRaises(ValueError):
                            main.rebuild_censo_from_dataframe(censo, frame)
                    main.db.session.rollback()
                    self.assertEqual(main.CensoRecord.query.filter_by(censo_id=ids[0]).count(), 1003)
                    self.assertEqual(main.db.session.get(main.Censo, ids[0]).total_records, 1003)
            finally:
                main.db.session.rollback()
                main.CensoRecord.query.filter_by(censo_id=ids[0]).delete()
                main.Censo.query.filter_by(id=ids[0]).delete()
                main.UploadedFile.query.filter_by(id=ids[1]).delete()
                main.db.session.commit()
