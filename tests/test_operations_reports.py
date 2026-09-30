import unittest
from datetime import date, time
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker

from tests.test_integration import app
from tests.test_form_structure import PageStructure
from gestion5s import web
from gestion5s.operations_reports import (SOURCES, SOURCE_MAP, build_operations_report, chart_data,
                                          parse_filters)
from gestion5s.operations_report_routes import BASE_PATH


class OperationsReportsTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        web.Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.session_patch = patch.object(web, "SessionLocal", self.sessions)
        self.session_patch.start()
        self.client = app.test_client()
        self.args = dict(start_date="2026-09-01", end_date="2026-09-03")

    def tearDown(self):
        self.session_patch.stop()
        self.engine.dispose()

    def create(self, entity, day=date(2026, 9, 1), **values):
        source, model = SOURCE_MAP[entity], web.ENTITY_MODEL[entity]
        values = {source["date"]: day, **values}
        if entity == "duplicidades":
            values.setdefault("semana", 36)
        elif entity == "robos":
            values.setdefault("hora", time(8, 15, 42))
        with self.sessions() as db:
            record = model(**values)
            db.add(record)
            db.commit()
            return record.id

    def seed(self):
        self.create("miscelaneo", estado="No iniciada", empresa="Ácme", ot="000123")
        self.create("miscelaneo", day=None, fecha_inicio=date(2026, 9, 2), estado="Aprobada", empresa=" ACME ")
        self.create("solicitud_ot", day=date(2026, 9, 3), estado="Completada", ot="OT0001")
        self.create("solicitud_ot", estado="Eliminado")
        self.create("solicitud_ot", day=None, estado="No iniciada")
        self.create("reclamos", estatus="Abierto")
        self.create("reclamos", day=date(2026, 9, 3), estatus="Cerrado")
        self.create("reclamos", day=date(2026, 9, 2), estatus="")
        self.create("robos", day=date(2026, 9, 2), recepciona="Recepción")
        self.create("duplicidades", day=date(2026, 9, 3), estatus="Abierto")
        self.create("habitaciones_bloqueadas", fecha_liberada_facility=date(2026, 9, 2))
        self.create("habitaciones_bloqueadas", day=None)
        self.create("desviaciones")
        self.create("desviaciones", day=date(2026, 9, 3), acciones="Acción correctiva")
        self.create("solicitud_ot", day=date(2026, 8, 31), fecha_inicio=date(2026, 9, 1), estado="Abierto")
        self.create("desviaciones", day=date(2026, 9, 4))
        with self.sessions() as db:
            for entity in ("samtech_qr", "samtech_usuarios"):
                db.add(web.ENTITY_MODEL[entity](fecha_creacion=date(2026, 9, 1), ticket="000123", estado="Completada"))
            db.commit()

    def result(self, export=None, **args):
        filters = parse_filters({**self.args, **args})
        with self.sessions() as db:
            return build_operations_report(db, web.ENTITY_MODEL, filters, export=export)

    def export(self, category="general", **args):
        url = BASE_PATH + (".xlsx" if category == "general" else f"/{category}.xlsx")
        response = self.client.get(url, query_string={**self.args, **args})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", response.content_type)
        return response.data

    def test_navigation_all_tabs_default_dates_and_forms(self):
        for url in ("/", "/gestion-5s/panel", "/reports/usuarios", BASE_PATH):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertIn(BASE_PATH, PageStructure(response.get_data(as_text=True)).links)
        for source in SOURCES:
            response = self.client.get(BASE_PATH, query_string={**self.args, "tab": source["entity"]})
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            self.assertIn("Descargar esta pestaña", html)
            self.assertIn("op-detail-scroll", html)
            for field in web.list_fields(source["entity"], web.ENTITY_MODEL[source["entity"]]()):
                self.assertIn(field["label"], html)
        default = parse_filters({})
        self.assertEqual(default["end"], date.today())
        self.assertEqual(default["start"], date.today().replace(day=1))

    def test_counts_seven_disjoint_sources_boundaries_and_current_states(self):
        self.seed()
        report = self.result()
        self.assertEqual([c["entity"] for c in report["categories"]], list(SOURCE_MAP))
        self.assertEqual(report["daily"], [5, 3, 4])
        self.assertEqual({k: report["totals"][k] for k in ("records", "open", "closed", "unknown", "not_applicable")},
                         dict(records=12, open=3, closed=3, unknown=2, not_applicable=4))
        self.assertEqual(report["totals"]["rate"], .5)
        self.assertEqual(report["totals"]["undated_available"], 2)
        self.assertEqual(report["totals"]["undated"], 0)
        by_entity = {c["entity"]: c for c in report["categories"]}
        self.assertEqual(by_entity["miscelaneo"]["totals"]["companies"], 1)
        self.assertEqual(by_entity["miscelaneo"]["companies"][0]["count"], 2)
        self.assertEqual(by_entity["desviaciones"]["totals"]["tracked"], 1)
        self.assertEqual(by_entity["desviaciones"]["totals"]["untracked"], 1)
        self.assertEqual(by_entity["habitaciones_bloqueadas"]["totals"]["tracked"], 1)
        for entity in ("robos", "habitaciones_bloqueadas", "desviaciones"):
            self.assertIsNone(by_entity[entity]["totals"]["rate"])
            self.assertEqual(by_entity[entity]["totals"]["closed"], 0)
        with self.sessions() as db:
            self.assertEqual(db.query(web.SamtechUsuarioEntry).count(), 1)
            self.assertEqual(db.query(web.ENTITY_MODEL["samtech_qr"]).count(), 1)

    def test_optional_dates_in_totals_and_exports_without_fake_trend_dates(self):
        self.seed()
        report = self.result(export="general", include_undated="1")
        self.assertEqual(report["daily"], [5, 3, 4])
        self.assertEqual(report["totals"]["records"], 14)
        self.assertEqual(report["totals"]["undated"], 2)
        self.assertEqual(report["totals"]["open"], 4)
        self.assertAlmostEqual(report["totals"]["rate"], 3 / 7)
        self.assertEqual(sum(len(c["details"]) for c in report["categories"]), 14)
        book = load_workbook(BytesIO(self.export(include_undated="1")), data_only=True)
        self.assertEqual(book["Evolución diaria"]["A10"].value, "Sin fecha")
        self.assertEqual(book["Evolución diaria"]["I10"].value, 2)
        self.assertEqual(book["Evolución diaria"]["I11"].value, 14)
        ot = book["Solicitudes y OT de usuario"]
        self.assertIsNone(ot.cell(9, 30).value)  # La tercera fila de detalle no tiene fecha.

    def test_search_same_scope_for_cards_details_and_exports_and_literal_symbols(self):
        self.seed()
        self.create("solicitud_ot", comentario="100%_listo", ot="001234")
        for q, expected in (("accion", 1), ("acme", 2), ("%_", 1), ("001234", 1), ("sin coincidencia", 0)):
            with self.subTest(q=q):
                report = self.result(export="general", q=q)
                self.assertEqual(report["totals"]["records"], expected)
                self.assertEqual(sum(len(c["details"]) for c in report["categories"]), expected)
                book = load_workbook(BytesIO(self.export(q=q)), data_only=True)
                self.assertEqual(book["Resumen general"]["B16"].value, expected)
        html = self.client.get(BASE_PATH, query_string={**self.args, "q": '<script>alert(1)</script>'}).get_data(as_text=True)
        self.assertNotIn('<script>alert(1)</script>', html)

    def test_pagination_full_exports_and_full_business_columns(self):
        for index in range(57):
            self.create("solicitud_ot", ot=f"{index:06d}", observacion="Dato histórico", comentario="Comentario", estado="Completada")
        first = next(c for c in self.result(tab="solicitud_ot")["categories"] if c["entity"] == "solicitud_ot")
        second = next(c for c in self.result(tab="solicitud_ot", page="2")["categories"] if c["entity"] == "solicitud_ot")
        self.assertEqual(len(first["details"]), 50)
        self.assertEqual(len(second["details"]), 7)
        self.assertEqual(first["totals"]["records"], second["totals"]["records"])
        self.assertEqual((second["first_row"], second["last_row"], second["pages"]), (51, 57, 2))
        self.assertTrue(set(r["ot"] for r in first["details"]).isdisjoint(r["ot"] for r in second["details"]))
        book = load_workbook(BytesIO(self.export("solicitud_ot", page="2")), data_only=True)
        self.assertEqual(book.sheetnames, ["Resumen", "Evolución diaria", "Solicitudes y OT de usuario"])
        detail = book["Solicitudes y OT de usuario"]
        self.assertEqual(detail.max_row, 63)
        self.assertEqual(detail.max_column, 31)
        self.assertEqual(detail["A7"].value, "000056")
        self.assertEqual(detail["AC7"].value, "Dato histórico")
        self.assertEqual(detail.freeze_panes, "A7")

    def test_general_and_every_individual_workbook_reconcile_and_preserve_types(self):
        self.seed()
        self.create("reclamos", n_solicitud="0009", descripcion_problema='=HYPERLINK("https://example.invalid","x")',
                    plan_accion="Línea 1; con ñ\nLínea 2", ingresar_contacto="00123456789")
        blob = self.export()
        cached = load_workbook(BytesIO(blob), data_only=True)
        formulas = load_workbook(BytesIO(blob), data_only=False)
        self.assertEqual(cached.sheetnames, ["Resumen general", "Evolución diaria", *[c["sheet"] for c in SOURCES]])
        self.assertEqual(cached["Resumen general"]["B16"].value, 13)
        self.assertEqual(cached["Evolución diaria"]["I10"].value, 13)
        self.assertEqual(formulas["Resumen general"]["G16"].data_type, "f")
        errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!"}
        self.assertFalse(any(c.value in errors for sheet in cached for row in sheet for c in row if isinstance(c.value, str)))
        self.assertEqual(formulas["Reclamos de usuarios"]["E9"].data_type, "s")
        self.assertTrue(formulas["Reclamos de usuarios"]["E9"].value.startswith("=HYPERLINK"))
        self.assertEqual(cached["Misceláneos"]["A8"].value, "000123")
        self.assertEqual(cached["Robos y hurtos"]["B7"].value, time(8, 15, 42))
        self.assertEqual(cached["Desviaciones"]["N7"].value, "Acción correctiva")
        with ZipFile(BytesIO(blob)) as archive:
            self.assertEqual(len([n for n in archive.namelist() if n.startswith("xl/charts/chart") and n.endswith(".xml")]), 2)
        for category in SOURCES:
            with self.subTest(entity=category["entity"]):
                single = load_workbook(BytesIO(self.export(category["entity"])), data_only=True)
                self.assertEqual(single.sheetnames, ["Resumen", "Evolución diaria", category["sheet"]])
                self.assertEqual(single["Resumen"]["B9"].value, cached[category["sheet"]].max_row - 6)

    def test_empty_data_no_division_errors_and_no_invented_undated_rows(self):
        report = self.result()
        self.assertIsNone(report["totals"]["rate"])
        self.assertEqual(report["totals"]["records"], 0)
        book = load_workbook(BytesIO(self.export(include_undated="1")), data_only=True)
        self.assertEqual(book["Resumen general"]["B16"].value, 0)
        self.assertIn(book["Resumen general"]["G16"].value, (None, ""))
        self.assertEqual(book["Evolución diaria"]["I10"].value, 0)
        html = self.client.get(BASE_PATH, query_string=self.args).get_data(as_text=True)
        self.assertIn("No hay registros para estos filtros", html)
        self.assertNotIn('id="operationsTrend"', html)

    def test_invalid_parameters_errors_do_not_export_unfiltered_data(self):
        for extra in (dict(start_date=""), dict(end_date="2026-08-30"), dict(start_date="2026-02-30"),
                      dict(tab="anything"), dict(page="-1"), dict(page="1.5"), dict(q="x"*201), dict(include_undated="yes")):
            with self.subTest(extra=extra):
                for path in (BASE_PATH, BASE_PATH + ".xlsx"):
                    r = self.client.get(path, query_string={**self.args, **extra})
                    self.assertEqual(r.status_code, 400)
                    self.assertEqual(r.headers["Cache-Control"], "no-store")
        self.assertEqual(self.client.get(BASE_PATH + "/samtech_qr.xlsx", query_string=self.args).status_code, 404)
        with patch("gestion5s.operations_report_routes.load_report", side_effect=SQLAlchemyError("test")):
            for path in (BASE_PATH, BASE_PATH + ".xlsx"):
                response = self.client.get(path, query_string=self.args)
                self.assertEqual(response.status_code, 503)
                self.assertIn("Inténtalo nuevamente", response.get_data(as_text=True))

    def test_monthly_chart_aggregation_preserves_counts(self):
        self.seed()
        report = self.result(start_date="2026-01-01", end_date="2026-09-30")
        chart = chart_data(report)
        self.assertTrue(chart["monthly"])
        self.assertEqual(chart["labels"], [f"{month:02d}/2026" for month in range(1,10)])
        self.assertEqual(sum(sum(c["values"]) for c in chart["series"]), report["totals"]["records"])
        daily = chart_data(self.result())
        self.assertEqual(daily["labels"], ["01/09/2026", "02/09/2026", "03/09/2026"])


if __name__ == "__main__":
    unittest.main()
