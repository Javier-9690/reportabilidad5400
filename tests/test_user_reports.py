import json
import re
import unicodedata
import unittest
from datetime import date, datetime, timedelta
from io import BytesIO
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker

from tests.test_integration import FORM_DATA, app
from tests.test_form_structure import PageStructure
from gestion5s import web
from gestion5s.user_reports import SOURCES, build_user_report, classify_status


class UserReportsTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        web.Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.session_patch = patch.object(web, "SessionLocal", self.sessions)
        self.session_patch.start()
        self.client = app.test_client()
        self.filters = {"start_date": "2026-08-31", "end_date": "2026-09-12"}

    def tearDown(self):
        self.session_patch.stop()
        self.engine.dispose()

    def create(self, entity, day="2026-08-31", status="Cerrado", **extra):
        source = next(item for item in SOURCES if item["entity"] == entity)
        data = {**FORM_DATA[entity], source["date"]: day or ""}
        if source.get("date_fallback"):
            data[source["date_fallback"]] = ""
        if source["status"]:
            data[source["status"]] = status
        data.update(extra)
        response = self.client.post(f"/gestion-5s/panel?tab={entity}", data=data)
        self.assertEqual(response.status_code, 302, response.get_data(as_text=True))
        with self.sessions() as db:
            model = web.ENTITY_MODEL[entity]
            return db.query(model).order_by(model.id.desc()).first().id

    def result(self, **filters):
        params = {**self.filters, **filters}
        with self.sessions() as db:
            return build_user_report(db, web.ENTITY_MODEL, date.fromisoformat(params["start_date"]),
                                     date.fromisoformat(params["end_date"]))

    def export(self, **filters):
        response = self.client.get("/reports/usuarios.xlsx", query_string={**self.filters, **filters})
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True) if response.status_code != 200 else "")
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", response.content_type)
        self.assertIn("reporte_usuarios_", response.headers["Content-Disposition"])
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        return response.data

    def sample(self):
        self.create("duplicidades", status="Abierto")
        self.create("duplicidades", status="Cerrado")
        self.create("duplicidades", day="2026-09-01", status="")
        self.create("reclamos", status="Abierto")
        self.create("reclamos", status="Cerrado")
        self.create("solicitud_ot", status="No iniciada")
        self.create("samtech_usuarios", status="No iniciada", ticket="000001")
        self.create("samtech_usuarios", status="En progreso", ticket="000002")
        self.create("samtech_usuarios", status="Cerrado", ticket="000003")
        self.create("solicitud_ot", day="2026-09-01", status="En proceso")
        self.create("solicitud_ot", day="2026-09-01", status="Cerrado")
        self.create("samtech_usuarios", day="2026-09-01", status="")
        self.create("reclamos", day="2026-09-12", status="Cerrado")
        for _ in range(3):
            self.create("desviaciones")
        for _ in range(2):
            self.create("desviaciones", day="2026-09-12")
        for entity in ("duplicidades", "reclamos", "solicitud_ot", "samtech_usuarios", "desviaciones"):
            self.create(entity, day="2026-08-30")
            self.create(entity, day="2026-09-13")
        self.create("solicitud_ot", day=None)
        self.create("samtech_usuarios", day=None)

    def test_navigation_range_form_and_initial_state(self):
        for path in ("/", "/gestion-5s/panel", "/reports/usuarios"):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200)
            self.assertIn("/reports/usuarios", PageStructure(response.get_data(as_text=True)).links)
        response = self.client.get("/reports/usuarios")
        page = response.get_data(as_text=True)
        self.assertIn("Selecciona el rango de fechas", page)
        self.assertNotIn('id="exportUserReport"', page)
        self.assertNotIn('data-kpi=', page)
        self.assertIn('for="reportStart"', page)
        self.assertIn('for="reportEnd"', page)
        for path in ("/reports/dotacion-gerencia", "/reports/ocupabilidad", "/reports/egp", "/reports/fa"):
            self.assertEqual(self.client.get(path).status_code, 200)

    def test_daily_counts_include_five_sources_both_boundaries_and_empty_days(self):
        self.sample()
        # Otros módulos no contribuyen a este informe aunque tengan la misma fecha.
        self.assertEqual(self.client.post("/gestion-5s/panel?tab=miscelaneo", data=FORM_DATA["miscelaneo"]).status_code, 302)
        report = self.result()
        self.assertEqual(report["days"], [date(2026, 8, 31) + timedelta(days=i) for i in range(13)])
        self.assertEqual(report["counts"], {
            "records": [11, 4] + [0] * 10 + [3], "open": [5, 1] + [0] * 11,
            "closed": [3, 1] + [0] * 10 + [1], "unclassified": [3, 2] + [0] * 10 + [2],
        })
        totals = report["totals"]
        self.assertEqual({key: totals[key] for key in ("records", "open", "closed", "unclassified", "classified", "undated")},
                         {"records": 18, "open": 6, "closed": 5, "unclassified": 7, "classified": 11, "undated": 2})
        self.assertEqual((totals["without_status"], totals["unknown_status"]), (5, 2))
        self.assertAlmostEqual(totals["rate"], 5 / 11)
        self.assertEqual([c["totals"]["records"] for c in report["categories"]], [3, 3, 3, 4, 5])
        self.assertEqual([row["label"] for row in report["sections"][0]["rows"][:-1]],
                         ["Doble asignación", "Reclamos usuarios", "Solicitudes de usuarios", "Samtech usuarios", "Desviaciones clientes"])
        self.assertEqual([section["key"] for section in report["sections"]], ["records", "open", "closed", "unclassified", "comparison"])
        open_rows = report["sections"][1]["rows"]
        self.assertIn("solicitud_ot_en_progreso", [row["key"] for row in open_rows])
        self.assertEqual(sum(row["total"] for row in open_rows[:-1]), open_rows[-1]["total"])
        text = " ".join(report["conclusions"])
        self.assertIn("45,5%", text)
        self.assertIn("11 casos", text)
        self.assertIn("La categoría con más registros es Desviaciones clientes: 5 casos", text)
        self.assertIn("Revisar estados: 2 casos", text)
        self.assertIn("31/08/2026", text)
        self.assertIn("histórico", text)

    def test_status_normalization_and_unknown_states_never_become_closed(self):
        cases = {
            "  ABIERTA  ": "abierto", "No   Iniciado": "no_iniciada", "No-iniciada": "no_iniciada",
            "EN_PROGRESO": "en_progreso", "En proceso": "en_progreso", "en espera": "pendiente",
            unicodedata.normalize("NFD", "EN EJECUCIÓN"): "en_progreso",
            "CERRADO": "cerrado", "Resuelta": "cerrado", "Finalizada": "cerrado", "Completado": "cerrado",
            "No cerrado": "sin_clasificar", "Reabierto": "sin_clasificar", "Cancelado": "sin_clasificar",
            "Rechazado": "sin_clasificar", "": "sin_clasificar", "   ": "sin_clasificar",
        }
        for raw, expected in cases.items():
            with self.subTest(status=raw):
                self.assertEqual(classify_status(raw), expected)
                self.create("samtech_usuarios", status=raw)
        report = self.result()
        self.assertEqual(report["totals"]["closed"], 4)
        self.assertEqual(report["totals"]["unclassified"], 6)
        self.assertEqual(report["totals"]["open"], 7)
        with self.sessions() as db:
            self.assertIn("No cerrado", {row.estado for row in db.query(web.SamtechUsuarioEntry)})

    def test_closure_percentage_is_weighted_and_does_not_include_unclassified(self):
        for _ in range(9):
            self.create("reclamos", status="Abierto")
        self.create("reclamos")
        self.create("samtech_usuarios", day="2026-09-01")
        for _ in range(5):
            self.create("solicitud_ot", status="")
        report = self.result()
        rate = report["sections"][-1]["rows"][-1]
        self.assertEqual(rate["values"][:2], [0.1, 1])
        self.assertEqual(rate["values"][2:], [0] * 11)
        self.assertAlmostEqual(rate["total"], 2 / 11)
        self.assertNotAlmostEqual(rate["total"], (0.1 + 1) / 2)
        self.assertEqual(report["totals"]["records"], 16)

    def test_reference_dates_are_independent_from_end_dates_and_internal_creation(self):
        self.create("samtech_usuarios", day="2026-08-31", fecha_inicio="2026-08-20", fecha_termino="2026-10-15", fecha_aprobacion="2026-11-01")
        self.create("samtech_usuarios", day="2026-07-01", fecha_inicio="2026-08-31", fecha_termino="2026-09-01")
        self.create("solicitud_ot", day="2026-09-12")
        self.create("reclamos", day="2026-08-31")
        self.create("duplicidades", day="2026-08-31", fecha_cierre="2026-10-15")
        self.create("duplicidades", day="2026-07-01", fecha_cierre="2026-08-31")
        self.create("desviaciones", day="2026-09-12")
        report = self.result()
        self.assertEqual(report["counts"]["records"], [3] + [0] * 11 + [2])
        self.assertEqual(report["counts"]["closed"], [3] + [0] * 11 + [1])
        self.assertEqual(report["counts"]["unclassified"], [0] * 12 + [1])
        self.assertIn("estado actual", report["method_note"])

    def test_deviations_with_optional_blanks_count_without_inventing_a_state(self):
        for action in ("", "Cerrado", "Abierto", "Revisar habitación"):
            self.create("desviaciones", acciones=action, id_interno="00042")
        report = self.result()
        deviations = next(c for c in report["categories"] if c["entity"] == "desviaciones")
        self.assertEqual(deviations["totals"], {"records": 4, "open": 0, "closed": 0,
                                              "unclassified": 4, "rate": None, "without_status": 4})
        self.assertEqual(deviations["unknown_states"], [])
        self.assertEqual(report["totals"]["unknown_status"], 0)
        conclusions = " ".join(report["conclusions"])
        self.assertIn("Desviaciones clientes aporta 4 casos al total", conclusions)
        self.assertNotIn("Revisar estados", conclusions)
        for section in report["sections"][1:3]:
            self.assertFalse(any(row["key"].startswith("desviaciones_") for row in section["rows"]))
        page = self.client.get("/reports/usuarios", query_string=self.filters).get_data(as_text=True)
        self.assertIn("No aplica", page)
        self.assertNotIn("Estados sin clasificar dentro del rango", page)
        values = load_workbook(BytesIO(self.export()), data_only=True)
        sheet = values["Desviaciones clientes"]
        headers = [cell.value for cell in sheet[5]]
        actions = headers.index("Acciones")
        identifier = headers.index("ID")
        self.assertEqual([row[actions].value for row in list(sheet)[5:]], [None, "Cerrado", "Abierto", "Revisar habitación"])
        self.assertEqual([row[identifier].value for row in list(sheet)[5:]], ["00042"] * 4)
        self.assertEqual([row[-1].value for row in list(sheet)[5:]], ["Sin campo de estado"] * 4)
        self.assertEqual(values["Conclusiones"]["P15"].value, "No aplica")

    def test_html_dashboard_tables_chart_data_and_export_use_same_range(self):
        self.sample()
        response = self.client.get("/reports/usuarios", query_string=self.filters)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        page = response.get_data(as_text=True)
        self.assertIn('data-kpi="records">18<', page)
        self.assertIn('data-kpi="rate">45,5%<', page)
        self.assertIn('data-record-scroll-top', page)
        self.assertIn('data-report-row="total_records"', page)
        payload = json.loads(re.search(r'<script id="userReportChartData" type="application/json">(.*?)</script>', page, re.S).group(1))
        self.assertEqual(payload["categories"], ["Doble asignación", "Reclamos usuarios", "Solicitudes de usuarios", "Samtech usuarios", "Desviaciones clientes"])
        self.assertEqual(len(set(payload["colors"])), 5)
        self.assertEqual(payload["open"], [1, 1, 2, 2, 0])
        self.assertEqual(payload["closed"], [1, 2, 1, 1, 0])
        self.assertEqual(payload["unclassified"], [1, 0, 0, 1, 5])
        self.assertEqual(payload["dates"][0], "2026-08-31")
        self.assertEqual(payload["dates"][-1], "2026-09-12")
        self.assertEqual(sum(sum(values) for values in payload["daily"]), 18)
        links = PageStructure(page).links
        download = next(link for link in links if "/reports/usuarios.xlsx" in link)
        self.assertEqual(parse_qs(urlsplit(download).query), {key: [value] for key, value in self.filters.items()})
        for entity in ("duplicidades", "reclamos", "solicitud_ot", "samtech_usuarios", "desviaciones"):
            detail = next(link for link in links if f"vista={entity}" in link)
            self.assertEqual(parse_qs(urlsplit(detail).query), {"vista": [entity], "from": ["2026-08-31"], "to": ["2026-09-12"]})
            self.assertEqual(self.client.get(detail).status_code, 200)

    def test_excel_contains_formulas_cached_results_dates_charts_and_full_source_columns(self):
        self.sample()
        content = self.export()
        formulas, values = load_workbook(BytesIO(content)), load_workbook(BytesIO(content), data_only=True)
        self.assertEqual(values.sheetnames, ["Reporte diario", "Conclusiones", "Doble asignación", "Reclamos usuarios", "Solicitudes OT", "Samtech usuarios", "Desviaciones clientes"])
        daily = values["Reporte diario"]
        self.assertEqual(daily.max_column, 15)
        self.assertEqual([cell.value.date() for cell in daily[7][1:14]], [date(2026, 8, 31) + timedelta(days=i) for i in range(13)])
        self.assertEqual(daily.freeze_panes, "B8")
        self.assertEqual(daily["A7"].fill.fgColor.rgb, "FFF00000")
        total_row = next(row for row in daily if row[0].value == "Total registros")
        self.assertEqual([cell.value for cell in total_row[1:]], [11, 4] + [0] * 10 + [3, 18])
        self.assertEqual(total_row[1].fill.fgColor.rgb, "FFFFFF00")
        percent_row = next(row for row in daily if row[0].value == "% cerrado")
        self.assertAlmostEqual(percent_row[-1].value, 5 / 11)
        self.assertEqual(percent_row[-1].number_format, "0%")
        self.assertIn("COUNTIFS(", formulas["Reporte diario"]["B8"].value)
        self.assertIn("'Doble asignación'!", formulas["Reporte diario"]["B8"].value)
        self.assertIn("'Desviaciones clientes'!", formulas["Reporte diario"]["B12"].value)
        self.assertTrue(formulas["Reporte diario"][total_row[1].coordinate].value.startswith("=SUM("))
        conclusion = values["Conclusiones"]
        self.assertEqual([conclusion[cell].value for cell in ("A6", "E6", "I6", "M6")], [18, 6, 5, 7])
        self.assertAlmostEqual(conclusion["Q6"].value, 5 / 11)
        self.assertEqual(len(formulas["Conclusiones"]._charts), 2)
        for source, expected in zip(SOURCES, (3, 3, 3, 4, 5)):
            sheet = values[source["sheet"]]
            fields = web.list_fields(source["entity"], web.ENTITY_MODEL[source["entity"]]())
            extra = ["Fecha para reporte"] if source.get("report_date_field") else []
            self.assertEqual([cell.value for cell in sheet[5]], [field["label"] for field in fields] + extra + ["Estado agrupado"])
            self.assertEqual(sheet.max_row - 5, expected)
            self.assertEqual(sheet.freeze_panes, "A6")
            self.assertTrue(sheet.auto_filter.ref)
        self.assertEqual(values["Samtech usuarios"]["A6"].value, "000001")
        self.assertEqual(values["Samtech usuarios"]["A6"].data_type, "s")
        with ZipFile(BytesIO(content)) as archive:
            self.assertIsNone(archive.testzip())
            charts = [name for name in archive.namelist() if re.fullmatch(r"xl/charts/chart\d+.xml", name)]
            self.assertEqual(len(charts), 2)
            namespace = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
            for name in charts:
                root = ET.fromstring(archive.read(name))
                if root.find(".//c:lineChart", namespace) is not None:
                    series = root.findall(".//c:lineChart/c:ser", namespace)
                    self.assertEqual([item.find("c:tx/c:v", namespace).text for item in series],
                                     [source["label"] for source in SOURCES])
                else:
                    category_refs = root.findall(".//c:barChart/c:ser/c:cat/c:strRef/c:f", namespace)
                    self.assertEqual(len(category_refs), 3)
                    self.assertEqual({item.text for item in category_refs}, {"Conclusiones!$A$11:$A$15"})
                cached = [float(item.text) for item in root.findall(".//c:numCache/c:pt/c:v", namespace)]
                self.assertTrue(cached)
                self.assertTrue(any(value > 0 for value in cached))
        for sheet in values:
            self.assertFalse(any(cell.data_type == "e" for row in sheet for cell in row))

    def test_no_records_and_only_unknown_states_do_not_claim_successful_closure(self):
        report = self.result()
        self.assertEqual(report["totals"]["rate"], 0)
        self.assertEqual(report["conclusions"], ["No hay registros con fecha dentro del rango seleccionado."])
        values = load_workbook(BytesIO(self.export()), data_only=True)
        self.assertEqual(values["Conclusiones"]["Q6"].value, 0)
        self.create("samtech_usuarios", status="")
        report = self.result()
        self.assertEqual(report["totals"]["records"], 1)
        self.assertEqual(report["totals"]["classified"], 0)
        self.assertIn("No hay estados clasificables", " ".join(report["conclusions"]))
        self.assertNotIn("Prioridad de seguimiento", " ".join(report["conclusions"]))

    def test_invalid_or_missing_ranges_are_not_silently_expanded(self):
        cases = ({}, {"start_date": ""}, {"start_date": "2026-09-01"}, {"end_date": "2026-09-01"},
                 {"start_date": "2026-02-30", "end_date": "2026-03-01"},
                 {"start_date": "2026-09-02", "end_date": "2026-09-01"},
                 {"start_date": "01/09/2026", "end_date": "2026-09-02"},
                 {"start_date": "1899-01-01", "end_date": "1900-01-01"},
                 {"start_date": "1900-01-01", "end_date": "9999-12-31"})
        for params in cases:
            with self.subTest(params=params):
                response = self.client.get("/reports/usuarios.xlsx", query_string=params)
                self.assertEqual(response.status_code, 400)
                self.assertNotIn('data-kpi=', response.get_data(as_text=True))
                if params:
                    self.assertEqual(self.client.get("/reports/usuarios", query_string=params).status_code, 400)

    def test_leap_days_year_boundaries_and_one_day_exports(self):
        for start, end, expected in (("2024-02-28", "2024-03-01", 3), ("2026-12-31", "2027-01-01", 2),
                                     ("2026-08-31", "2026-08-31", 1)):
            with self.subTest(start=start, end=end):
                self.create("samtech_usuarios", day=end)
                report = self.result(start_date=start, end_date=end)
                self.assertEqual(len(report["days"]), expected)
                self.assertEqual(report["counts"]["records"][-1], 1)
                values = load_workbook(BytesIO(self.export(start_date=start, end_date=end)), data_only=True)
                self.assertEqual(values["Reporte diario"].max_column, expected + 2)
                self.assertEqual(values["Conclusiones"]["A6"].value, 1)

    def test_html_escaping_literal_excel_strings_and_read_only_behavior(self):
        text = '<script>alert("x")</script>'
        rid = self.create("samtech_usuarios", status=text, ticket="=1+1", comentario="https://example.org", falla="@SUM(A1:A2)")
        response = self.client.get("/reports/usuarios", query_string=self.filters)
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(text, response.get_data(as_text=True))
        values = load_workbook(BytesIO(self.export()))
        sheet = values["Samtech usuarios"]
        self.assertEqual(sheet["A6"].value, "=1+1")
        self.assertEqual(sheet["A6"].data_type, "s")
        self.assertIsNone(sheet["O6"].hyperlink)
        self.assertEqual(sheet["N6"].value, text)
        with self.sessions() as db:
            record = db.get(web.SamtechUsuarioEntry, rid)
            self.assertEqual((record.ticket, record.estado), ("=1+1", text))
            self.assertEqual(db.query(web.SamtechUsuarioEntry).count(), 1)

    def test_reports_refresh_after_record_edits(self):
        rid = self.create("samtech_usuarios", status="No iniciada")
        self.assertEqual(self.result()["totals"]["open"], 1)
        with self.sessions() as db:
            db.get(web.SamtechUsuarioEntry, rid).estado = "Cerrado"
            db.commit()
        report = self.result()
        self.assertEqual((report["totals"]["open"], report["totals"]["closed"]), (0, 1))
        values = load_workbook(BytesIO(self.export()), data_only=True)
        self.assertEqual((values["Conclusiones"]["E6"].value, values["Conclusiones"]["I6"].value), (0, 1))

    def test_database_failures_and_excel_limits_show_errors_instead_of_empty_success(self):
        with patch.object(web, "SessionLocal", side_effect=SQLAlchemyError("test failure")), self.assertLogs(app.logger, level="ERROR"):
            for path in ("/reports/usuarios", "/reports/usuarios.xlsx"):
                response = self.client.get(path, query_string=self.filters)
                self.assertEqual(response.status_code, 503)
                self.assertNotIn('data-kpi=', response.get_data(as_text=True))
        self.create("samtech_usuarios", comentario="x" * 32768)
        response = self.client.get("/reports/usuarios.xlsx", query_string=self.filters)
        self.assertEqual(response.status_code, 400)
        self.assertIn("32.767 caracteres", response.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
