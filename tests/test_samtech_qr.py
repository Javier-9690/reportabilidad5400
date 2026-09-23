"""Aislamiento de las cargas QR y general, y trazabilidad del reporte."""

import csv
import json
import re
import tempfile
import unittest
from datetime import date
from io import BytesIO, StringIO
from unittest.mock import patch
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy import create_engine, event, text
from sqlalchemy.orm import sessionmaker

from tests import test_order_import as helpers
from tests.test_integration import FORM_DATA
from tests.test_edit_records import FormValues
from tests.test_form_structure import PageStructure
from gestion5s import web
from gestion5s.editing import SAMTECH_USER_FIELDS
from gestion5s.orders import ORDER_IMPORT_ENTITIES, parse_orders


class SamtechQRTest(unittest.TestCase):
    setUp = helpers.OrderImportTest.setUp
    tearDown = helpers.OrderImportTest.tearDown
    csrf = helpers.OrderImportTest.csrf
    confirm = helpers.OrderImportTest.confirm
    rows = helpers.OrderImportTest.rows
    report = helpers.OrderImportTest.report

    def preview(self, rows, entity="samtech_qr", **kwargs):
        return helpers.OrderImportTest.preview(self, rows, entity=entity, **kwargs)

    def snapshot(self):
        return {key: self.rows(key) for key in (*ORDER_IMPORT_ENTITIES, "samtech_usuarios", "reclamos")}

    def create(self, **values):
        response = self.client.post("/gestion-5s/panel?tab=samtech_qr", data={**FORM_DATA["samtech_qr"], **values})
        self.assertEqual(response.status_code, 302)
        return self.rows("samtech_qr")[-1]

    def test_template_form_and_optional_dates_preserve_the_fifteen_fields(self):
        page = self.client.get("/gestion-5s/panel?tab=samtech_qr").get_data(as_text=True)
        self.assertIn("Solicitudes de usuarios Samtech QR", page)
        self.assertIn("/gestion-5s/registros?vista=samtech_qr", PageStructure(page).links)
        self.assertEqual(PageStructure(page).nested_forms, 0)
        response = self.client.get("/gestion-5s/template/samtech_qr.xlsx")
        workbook = load_workbook(BytesIO(response.data))
        self.assertEqual([cell.value for cell in workbook.active[1]], [label for _, label in SAMTECH_USER_FIELDS])
        self.assertEqual(workbook.active["A2"].number_format, "@")
        for column in range(10, 14):
            self.assertEqual(workbook.active.cell(2, column).number_format, "DD/MM/YYYY")
        record = self.create(ticket="0000007", fecha_creacion="", comentario="Revisar instalación\nSegundo piso")
        self.assertEqual(record["ticket"], "0000007")
        self.assertTrue(all(record[name] is None for name, _ in SAMTECH_USER_FIELDS if name.startswith("fecha_")))
        for data in ({}, {"ticket": "x", "fecha_creacion": "2026-02-30"}):
            self.assertEqual(self.client.post("/gestion-5s/panel?tab=samtech_qr", data=data).status_code, 422)
        self.assertEqual(len(self.rows("samtech_qr")), 1)

    def test_qr_preview_confirmation_keeps_all_specialties_and_preserves_every_other_base(self):
        self.create(ticket="ANTERIOR")
        before = self.snapshot()
        orders = [helpers.order("001", "CARPINTERIA MENOR", estado="Aprobada"),
                  helpers.order("002", "CARPINTERIA"), helpers.order("003", "")]
        preview = self.preview(orders)
        page = preview.get_data(as_text=True)
        self.assertIn("Confirmar y reemplazar solicitudes QR", page)
        self.assertNotIn("Confirmar y reemplazar ambas bases", page)
        self.assertIn("CARPINTERIA MENOR", page)
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(self.confirm(preview, accepted=False).status_code, 400)
        self.assertEqual(self.snapshot(), before)
        # Un destino enviado por el navegador no cambia el alcance firmado.
        payload = FormValues(page).values
        payload.update(confirm_replace="1", entity="miscelaneo", return_tab="solicitud_ot")
        result = self.client.post("/gestion-5s/import/ordenes/confirm", data=payload)
        self.assertEqual(result.location, "/gestion-5s/registros?vista=samtech_qr")
        self.assertEqual([row["ticket"] for row in self.rows("samtech_qr")], ["001", "002", "003"])
        for entity, rows in before.items():
            if entity != "samtech_qr":
                self.assertEqual(self.rows(entity), rows)
        self.assertEqual(self.confirm(preview).status_code, 409)
        self.assertEqual(self.confirm(self.preview(orders)).status_code, 302)
        self.assertEqual(len(self.rows("samtech_qr")), 3)

    def test_general_and_qr_previews_can_be_confirmed_independently(self):
        qr = self.preview([helpers.order("QR", "CARPINTERIA MENOR")])
        general = self.preview([helpers.order("OT"), helpers.order("MISC", "CARPINTERIA MENOR")], entity="miscelaneo")
        self.assertEqual(self.confirm(general).status_code, 302)
        self.assertEqual(self.confirm(qr).status_code, 302)
        self.assertEqual([row["ticket"] for row in self.rows("samtech_qr")], ["QR"])
        self.assertEqual([row["ot"] for row in self.rows("solicitud_ot")], ["OT"])
        self.assertEqual([row["ot"] for row in self.rows("miscelaneo")], ["MISC"])
        before_qr = self.rows("samtech_qr")
        self.assertEqual(self.confirm(self.preview([helpers.order("OTHER")], entity="solicitud_ot")).status_code, 302)
        self.assertEqual(self.rows("samtech_qr"), before_qr)

    def test_cancel_invalid_file_and_database_failure_never_erase_qr_records(self):
        self.create(ticket="CONSERVAR")
        before = self.snapshot()
        preview = self.preview([helpers.order()])
        result = self.client.post("/gestion-5s/import/ordenes/cancel", data=FormValues(preview.get_data(as_text=True)).values)
        self.assertEqual(result.location, "/gestion-5s/panel?tab=samtech_qr")
        self.assertEqual(self.confirm(preview).status_code, 409)
        for rows in ([], [helpers.order(), helpers.order(fecha_termino="32/12/2026")],
                     [helpers.order(estado="x" * 101)]):
            self.preview(rows, expected=400)
            self.assertEqual(self.snapshot(), before)
        preview = self.preview([helpers.order()])
        def fail(db, models, groups):
            db.execute(models["samtech_qr"].__table__.insert(), groups["samtech_qr"])
            raise SQLAlchemyError("fallo simulado después del borrado y la inserción")
        with patch("gestion5s.order_import_routes._insert_groups", side_effect=fail), self.assertLogs(web.app.logger, level="ERROR"):
            self.assertEqual(self.confirm(preview).status_code, 503)
        self.assertEqual(self.snapshot(), before)

    def test_edit_after_preview_requires_new_confirmation_without_losing_changes(self):
        self.create(ticket="CONSERVAR")
        preview = self.preview([helpers.order()])
        with self.sessions() as db:
            db.query(web.SamtechQRUsuarioEntry).first().comentario = "Edición posterior"
            db.commit()
        changed = self.snapshot()
        self.assertEqual(self.confirm(preview).status_code, 409)
        self.assertEqual(self.snapshot(), changed)

    def test_blank_cells_reordered_headers_and_creation_fallback_match_search_export_and_dashboard(self):
        orders = [helpers.order("001", "CARPINTERIA MENOR", fecha_creacion=None, fecha_inicio="12/09/2026",
                                estado="Aprobada", falla="Revisión <urgente>\n" + "x" * 500),
                  helpers.order("002", fecha_creacion=None, fecha_inicio=None), [None] * 15]
        content = helpers.workbook_bytes(orders, reverse=True)
        parsed = parse_orders(content, entity="samtech_qr")
        self.assertEqual(parsed["total"], 2)
        self.assertEqual(parsed["fallback"], {"samtech_qr": 1})
        self.assertEqual(parsed["undated"], {"samtech_qr": 1})
        self.assertEqual(self.confirm(self.preview([], content=content)).status_code, 302)
        query = {"from": "2026-09-12", "to": "2026-09-12", "q": "revision"}
        response = self.client.get("/gestion-5s/registros", query_string={"vista": "samtech_qr", **query})
        page = response.get_data(as_text=True)
        self.assertIn("Revisión &lt;urgente&gt;", page)
        self.assertIn("Fecha creación (Fecha inicio si falta)", page)
        exported = self.client.get("/gestion-5s/download/samtech_qr.csv", query_string=query)
        rows = list(csv.DictReader(StringIO(exported.data.decode("utf-8-sig"))))
        self.assertEqual([row["Ticket"] for row in rows], ["001"])
        self.assertEqual(rows[0]["Fecha creación"], "")
        self.assertEqual(rows[0]["Fecha inicio"], "2026-09-12")
        self.assertEqual(rows[0]["Falla"], orders[0]["falla"])
        dashboard = self.client.get("/gestion-5s/dashboard", query_string=query).get_data(as_text=True)
        self.assertIn('id="samtechQRChart"', dashboard)
        series = json.loads(re.search(r"const series = (.*?);", dashboard).group(1))
        self.assertEqual(series["samtech_qr"], [1])

    def test_report_sources_and_excel_use_qr_separately_and_combine_both_general_categories(self):
        general = [helpers.order("001", "CARPINTERIA MENOR", estado="Aprobada"),
                   helpers.order("002", estado="En Progreso"),
                   helpers.order("003", estado="Eliminado", fecha_creacion=None, fecha_inicio="12/09/2026")]
        qr = [helpers.order("001", "CARPINTERIA MENOR", estado="Aprobada"),
              helpers.order("002", estado="No Iniciada"),
              helpers.order("004", estado="Aprobada", fecha_creacion=None, fecha_inicio=None)]
        self.assertEqual(self.confirm(self.preview(general, entity="solicitud_ot")).status_code, 302)
        self.assertEqual(self.confirm(self.preview(qr)).status_code, 302)
        report = self.report(details=True)
        categories = {item["label"]: item for item in report["categories"]}
        users, totals = categories["Solicitudes de usuarios"], categories["Solicitudes totales Samtech"]
        self.assertEqual(users["entity"], "samtech_qr")
        self.assertEqual([users["totals"][key] for key in ("records", "open", "closed")], [2, 1, 1])
        self.assertEqual(users["undated"], 1)
        self.assertEqual([totals["totals"][key] for key in ("records", "open", "closed", "unclassified")], [3, 1, 1, 1])
        self.assertEqual({row["ot"] for row in totals["details"]}, {"001", "002", "003"})
        self.assertEqual({row["_origen"] for row in totals["details"]}, {"Misceláneos", "Solicitudes OT"})
        self.assertEqual(self.report()["counts"], report["counts"])
        filters = {"start_date": "2026-08-31", "end_date": "2026-09-12"}
        page = self.client.get("/reports/usuarios", query_string=filters).get_data(as_text=True)
        self.assertIn("no representa tickets únicos", page)
        for entity in ("miscelaneo", "solicitud_ot", "samtech_qr"):
            self.assertIn(f"vista={entity}&", page)
        response = self.client.get("/reports/usuarios.xlsx", query_string=filters)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), data_only=True)
        self.assertNotIn("Samtech usuarios", workbook.sheetnames)
        for sheet, tickets in (("Solicitudes usuarios QR", {"001", "002"}),
                               ("Solicitudes totales Samtech", {"001", "002", "003"})):
            self.assertEqual({row[0] for row in workbook[sheet].iter_rows(min_row=6, values_only=True)}, tickets)
        rows = {}
        for row in workbook["Reporte diario"]:
            rows.setdefault(row[0].value, row)
        self.assertEqual(rows["Solicitudes de usuarios"][-1].value, 2)
        self.assertEqual(rows["Solicitudes totales Samtech"][-1].value, 3)
        self.assertEqual(rows["Total registros (suma de categorías)"][-1].value, report["totals"]["records"])

    def test_qr_pagination_search_and_delete_all_cover_the_full_filtered_result(self):
        self.assertEqual(self.confirm(self.preview([helpers.order(f"PAGE-{i:03}") for i in range(105)])).status_code, 302)
        other = self.create(ticket="OTRO", fecha_creacion="2026-08-31")
        general_before = {key: self.rows(key) for key in ("miscelaneo", "solicitud_ot")}
        query = {"vista": "samtech_qr", "q": "PAGE", "page": 2}
        page = self.client.get("/gestion-5s/registros", query_string=query).get_data(as_text=True)
        self.assertIn("101–105 de 105", page)
        self.assertIn("Eliminar todos (105)", page)
        csv_data = self.client.get("/gestion-5s/download/samtech_qr.csv", query_string=query)
        self.assertEqual(len(list(csv.DictReader(StringIO(csv_data.data.decode("utf-8-sig"))))), 105)
        preview = self.client.post("/gestion-5s/delete/samtech_qr/bulk/confirm", data={
            "csrf_token": FormValues(page).values["csrf_token"], "mode": "all", "q": "PAGE",
        })
        self.assertEqual(preview.status_code, 200)
        response = self.client.post("/gestion-5s/delete/samtech_qr/bulk", data=FormValues(preview.get_data(as_text=True)).values)
        self.assertEqual(response.status_code, 302)
        self.assertEqual([r["id"] for r in self.rows("samtech_qr")], [other["id"]])
        for key, rows in general_before.items():
            self.assertEqual(self.rows(key), rows)

    def test_report_does_not_mix_import_versions_when_another_connection_writes(self):
        # WAL permite confirmar una carga durante la lectura, como ocurre con
        # distintos workers de producción. El reporte debe conservar su versión.
        with tempfile.TemporaryDirectory() as folder:
            engine = create_engine("sqlite:///" + str(Path(folder) / "concurrent.db"))
            with engine.connect() as conn:
                conn.execute(text("PRAGMA journal_mode=WAL"))
            web.Base.metadata.create_all(engine)
            sessions = sessionmaker(bind=engine)
            with sessions() as db:
                db.add(web.MiscelaneoEntry(ot="OLD-M", fecha_creacion=date(2026, 8, 31)))
                db.add(web.SolicitudOTEntry(ot="OLD-OT", fecha_creacion=date(2026, 8, 31)))
                db.commit()
            changed = []
            def import_during_report(conn, cursor, statement, parameters, context, many):
                if changed or not statement.lstrip().upper().startswith("SELECT"):
                    return
                changed.append(True)
                with sessions() as writer:
                    writer.add(web.MiscelaneoEntry(ot="NEW-M", fecha_creacion=date(2026, 8, 31)))
                    writer.add(web.SolicitudOTEntry(ot="NEW-OT", fecha_creacion=date(2026, 8, 31)))
                    writer.commit()
            event.listen(engine, "after_cursor_execute", import_during_report)
            with patch.object(web, "SessionLocal", sessions):
                query = {"start_date": "2026-08-31", "end_date": "2026-08-31"}
                first = self.client.get("/reports/usuarios", query_string=query)
                second = self.client.get("/reports/usuarios", query_string=query)
            event.remove(engine, "after_cursor_execute", import_during_report)
            for response, expected in ((first, 2), (second, 4)):
                self.assertEqual(response.status_code, 200)
                self.assertIn(f'data-kpi="records">{expected}<', response.get_data(as_text=True))
            self.assertTrue(changed)
            engine.dispose()


if __name__ == "__main__":
    unittest.main()
