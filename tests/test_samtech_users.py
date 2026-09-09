import csv
import html
import json
import re
import unittest
from datetime import date, datetime
from io import BytesIO, StringIO
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from tests.test_integration import app
from tests.test_edit_records import FormValues
from tests.test_form_structure import PageStructure
from gestion5s import web


HEADERS = [
    "Ticket", "División", "Área", "Lugar", "Ubicación", "Disciplina", "Especialidad", "Falla", "Empresa",
    "Fecha creación", "Fecha inicio", "Fecha término", "Fecha aprobación", "Estado", "Comentario",
]
NAMES = [
    "ticket", "division", "area", "lugar", "ubicacion", "disciplina", "especialidad", "falla", "empresa",
    "fecha_creacion", "fecha_inicio", "fecha_termino", "fecha_aprobacion", "estado", "comentario",
]
FORM = dict(zip(NAMES, [
    "000123", "Campamento", "Hotelería", "Pabellón A", "Habitación 007", "Mantención", "Electricidad",
    "Revisión <solicitada>.\nComprobar todas las conexiones.", "Empresa & Servicios",
    "2026-09-09", "2026-09-10", "2026-09-11", "2026-09-12", "En proceso", "Seguimiento & coordinación.\nÚltima línea.",
]))
DATES = NAMES[9:13]


class SamtechUsersTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        web.Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.session_patch = patch.object(web, "SessionLocal", self.sessions)
        self.session_patch.start()
        self.client = app.test_client()

    def tearDown(self):
        self.session_patch.stop()
        self.engine.dispose()

    def records(self):
        with self.sessions() as db:
            return [{column.name: getattr(record, column.name) for column in record.__table__.columns}
                    for record in db.query(web.SamtechUsuarioEntry).order_by(web.SamtechUsuarioEntry.id).all()]

    def create(self, data):
        response = self.client.post("/gestion-5s/panel?tab=samtech_usuarios", data=data)
        self.assertEqual(response.status_code, 302, response.get_data(as_text=True))
        return self.records()[-1]

    def template(self):
        response = self.client.get("/gestion-5s/template/samtech_usuarios.xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data))
        self.assertEqual([cell.value for cell in workbook.active[1]], HEADERS)
        self.assertEqual(workbook.active.freeze_panes, "A2")
        self.assertEqual(workbook.active["A2"].number_format, "@")
        for column in range(10, 14):
            self.assertEqual(workbook.active.cell(2, column).number_format, "DD/MM/YYYY")
        return workbook

    def upload(self, workbook):
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        response = self.client.post("/gestion-5s/import/samtech_usuarios",
                                    data={"file": (payload, "samtech_usuarios.xlsx")},
                                    content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        return response.get_data(as_text=True)

    def exported(self, **filters):
        response = self.client.get("/gestion-5s/download/samtech_usuarios.csv", query_string=filters)
        self.assertEqual(response.status_code, 200)
        reader = csv.DictReader(StringIO(response.data.decode("utf-8-sig")))
        self.assertEqual(reader.fieldnames, HEADERS)
        return list(reader)

    def assert_labels(self, page):
        labels = re.findall(r'<label[^>]*class="form-label[^\"]*"[^>]*>(.*?)</label>', page, re.S)
        self.assertEqual([html.unescape(label.strip()) for label in labels], HEADERS)
        self.assertEqual([name for name in FormValues(page).values if name in NAMES], NAMES)

    def test_all_fifteen_columns_in_forms_template_list_export_and_editor(self):
        self.assert_labels(self.client.get("/gestion-5s/panel?tab=samtech_usuarios").get_data(as_text=True))
        before = self.create(FORM)
        self.assertEqual({name: str(before[name]) for name in NAMES}, FORM)
        page = self.client.get("/gestion-5s/registros?vista=samtech_usuarios").get_data(as_text=True)
        self.assertEqual(re.findall(r"<th>(.*?)</th>", page), ["Acciones", *HEADERS, "Registro"])
        self.assertEqual(PageStructure(page).cells, FORM)
        self.assertIn('data-record-scroll-top', page)
        self.assertIn("Eliminar seleccionados", page)
        self.assertIn("Eliminar todos", page)
        self.assertEqual(self.exported(), [dict(zip(HEADERS, FORM.values()))])

        route = f"/gestion-5s/edit/samtech_usuarios/{before['id']}"
        page = self.client.get(route).get_data(as_text=True)
        self.assert_labels(page)
        data = FormValues(page).values
        changed = dict(FORM, ticket="000124", fecha_aprobacion="", estado="Resuelto")
        data.update(changed)
        self.assertEqual(self.client.post(route, data=data, follow_redirects=True).status_code, 200)
        after = self.records()[0]
        self.assertEqual((after["id"], after["creado"]), (before["id"], before["creado"]))
        self.assertIsNone(after["fecha_aprobacion"])
        self.assertEqual(self.exported(), [dict(zip(HEADERS, changed.values()))])

        workbook = self.template()
        workbook.epoch = CALENDAR_MAC_1904
        imported = dict(FORM, fecha_creacion=datetime(2026, 9, 9), fecha_inicio=date(2026, 9, 10),
                        fecha_termino=to_excel(datetime(2026, 9, 11), epoch=workbook.epoch),
                        fecha_aprobacion="12/09/2026")
        for column, name in enumerate(NAMES, 1):
            workbook.active.cell(2, column).value = imported[name]
        workbook.active["L2"].number_format = "General"
        self.assertIn("OK: 1 filas.", self.upload(workbook))
        self.assertEqual({name: str(self.records()[-1][name]) for name in NAMES}, FORM)

    def test_each_field_can_import_alone_including_all_four_blank_dates(self):
        workbook = self.template()
        for row, (name, value) in enumerate(FORM.items(), 3):
            for column, other in enumerate(NAMES, 1):
                workbook.active.cell(row, column).value = value if name == other else ("  " if column % 2 else None)
        self.assertIn("OK: 15 filas.", self.upload(workbook))
        self.assertEqual(len(self.records()), 15)
        for name, record in zip(NAMES, self.records()):
            self.assertEqual(str(record[name]), FORM[name])
            self.assertTrue(all(record[other] is None for other in NAMES if other != name))
        self.assertEqual(len(self.exported()), 15)
        self.assertEqual(len(self.exported(**{"from": "2026-09-09", "to": "2026-09-09"})), 1)
        record = self.create({"ticket": "000000"})
        self.assertTrue(all(record[name] is None for name in DATES))
        self.assertEqual(self.client.post("/gestion-5s/panel?tab=samtech_usuarios", data={}).status_code, 422)

    def test_invalid_dates_report_errors_without_saving_or_partially_importing(self):
        before = self.create(FORM)
        for name in DATES:
            with self.subTest(field=name):
                invalid = dict(FORM, **{name: "2026-02-30"})
                response = self.client.post("/gestion-5s/panel?tab=samtech_usuarios", data=invalid)
                self.assertEqual(response.status_code, 422)
                self.assertEqual(FormValues(response.get_data(as_text=True)).values[name], invalid[name])
                route = f"/gestion-5s/edit/samtech_usuarios/{before['id']}"
                data = FormValues(self.client.get(route).get_data(as_text=True)).values
                data.update(invalid)
                self.assertEqual(self.client.post(route, data=data).status_code, 422)
                workbook = Workbook()
                workbook.active.append(HEADERS)
                workbook.active.append(list(FORM.values()))
                workbook.active.append(list(invalid.values()))
                self.assertIn("Error importando samtech_usuarios: Fila 3:", self.upload(workbook))
                self.assertEqual(self.records(), [before])

    def test_dashboard_counts_only_creation_dates_and_keeps_samtech_separate(self):
        self.create({"ticket": "Sin creación", "fecha_inicio": "2026-09-09", "fecha_aprobacion": "2026-09-10"})
        self.assertNotIn("const series =", self.client.get("/gestion-5s/dashboard").get_data(as_text=True))
        for day in ("2026-08-31", "2026-09-09", "2026-09-09", "2026-09-10"):
            self.create({"ticket": "000007", "fecha_creacion": day, "fecha_inicio": "2026-09-09",
                         "fecha_termino": "2026-09-15", "fecha_aprobacion": "2026-09-16"})
        self.assertEqual(self.client.post("/gestion-5s/panel?tab=miscelaneo", data={
            "ot": "000007", "fecha_creacion": "2026-09-09",
        }).status_code, 302)
        for filters in ({"from": "2026-09-09", "to": "2026-09-11"}, {"semana": 89}):
            with self.subTest(filters=filters):
                response = self.client.get("/gestion-5s/dashboard", query_string=filters)
                self.assertEqual(response.status_code, 200)
                page = response.get_data(as_text=True)
                labels = json.loads(re.search(r"const labels = (.*?);", page).group(1))
                series = json.loads(re.search(r"const series = (.*?);", page).group(1))
                self.assertEqual(labels, ["2026-09-09", "2026-09-10"])
                self.assertEqual(series["samtech_usuarios"], [2, 1])
                self.assertEqual(series["miscelaneo"], [1, 0])
                self.assertNotIn("salidas", series)
                self.assertIn('id="samtechUsuariosChart"', page)
                self.assertRegex(page, r'<div class="number[^\"]*">3</div>\s*<div class="label">Samtech usuarios</div>')
                self.assertEqual(len(self.exported(**filters)), 3)
                self.assertEqual(len(self.exported()), 5)


if __name__ == "__main__":
    unittest.main()
