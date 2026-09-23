import csv
from datetime import date, datetime, timedelta
from io import BytesIO, StringIO
import json
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
from sqlalchemy import Column, MetaData, Table, create_engine, inspect, text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker

from tests.test_integration import app, FORM_DATA
from tests.test_edit_records import FormValues
from tests.test_form_structure import PageStructure
from gestion5s import web
from gestion5s.editing import ORDER_IMPORT_FIELDS
from gestion5s.orders import ORDER_ENTITIES, parse_orders
from gestion5s.order_import_routes import orders_snapshot
from gestion5s.user_reports import build_user_report, classify_status


def workbook_bytes(rows, headers=None, reverse=False, epoch=None):
    workbook = Workbook()
    if epoch:
        workbook.epoch = epoch
    sheet = workbook.active
    sheet.title = "ReportePlanificacion"
    headers = headers or [label for _, label in ORDER_IMPORT_FIELDS]
    sheet.append(list(reversed(headers)) if reverse else headers)
    for row in rows:
        values = [row.get(name) for name, _ in ORDER_IMPORT_FIELDS] if isinstance(row, dict) else list(row)
        sheet.append(list(reversed(values)) if reverse else values)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def order(ticket="000001", specialty="GASFITER", **extra):
    return {"ot": ticket, "especialidad": specialty, "fecha_creacion": "31/08/2026",
            "estado": "No Iniciada", "comentario": "Revisar\nSegunda línea", **extra}


class OrderImportTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        web.Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.session_patch = patch.object(web, "SessionLocal", self.sessions)
        self.session_patch.start()
        self.client = app.test_client()
        for entity in (*ORDER_ENTITIES, "reclamos", "samtech_usuarios"):
            response = self.client.post(f"/gestion-5s/panel?tab={entity}", data=FORM_DATA[entity])
            self.assertEqual(response.status_code, 302)

    def tearDown(self):
        self.session_patch.stop()
        self.engine.dispose()

    def snapshot(self):
        with self.sessions() as db:
            return orders_snapshot(db)

    def csrf(self, client=None, entity="solicitud_ot"):
        client = client or self.client
        page = client.get(f"/gestion-5s/panel?tab={entity}").get_data(as_text=True)
        return FormValues(page).values["csrf_token"]

    def preview(self, rows, entity="solicitud_ot", content=None, expected=200):
        response = self.client.post(f"/gestion-5s/import/{entity}", data={
            "csrf_token": self.csrf(entity=entity),
            "file": (BytesIO(content if content is not None else workbook_bytes(rows)), "ordenes.xlsx"),
        })
        self.assertEqual(response.status_code, expected, response.get_data(as_text=True)[:800])
        return response

    def confirm(self, preview, accepted=True):
        data = FormValues(preview.get_data(as_text=True)).values
        if accepted:
            data["confirm_replace"] = "1"
        return self.client.post("/gestion-5s/import/ordenes/confirm", data=data)

    def rows(self, entity):
        with self.sessions() as db:
            return [{column.name: getattr(record, column.name) for column in record.__table__.columns}
                    for record in db.query(web.ENTITY_MODEL[entity]).order_by(web.ENTITY_MODEL[entity].id)]

    def report(self, details=False):
        with self.sessions() as db:
            return build_user_report(db, web.ENTITY_MODEL, date(2026, 8, 31), date(2026, 9, 12), include_details=details)

    def test_both_panels_and_templates_offer_the_same_confirmed_import(self):
        for entity in ORDER_ENTITIES:
            page = self.client.get(f"/gestion-5s/panel?tab={entity}").get_data(as_text=True)
            self.assertIn("CARPINTERIA MENOR", page)
            self.assertIn("Importar órdenes", page)
            self.assertTrue(FormValues(page).values["csrf_token"])
            self.assertEqual(PageStructure(page).nested_forms, 0)
            response = self.client.get(f"/gestion-5s/template/{entity}.xlsx")
            workbook = load_workbook(BytesIO(response.data))
            self.assertEqual([cell.value for cell in workbook.active[1]], [label for _, label in ORDER_IMPORT_FIELDS])

    def test_preview_and_cancel_leave_both_bases_intact(self):
        before = self.snapshot()
        preview = self.preview([order(), order("000002", "CARPINTERIA MENOR")], entity="miscelaneo")
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(preview.headers["Cache-Control"], "no-store")
        values = FormValues(preview.get_data(as_text=True)).values
        self.assertNotIn("confirm_replace", values)
        self.assertNotIn("comentario", values)  # Los datos completos no viajan en el formulario.
        self.assertEqual(self.confirm(preview, accepted=False).status_code, 400)
        self.assertEqual(self.snapshot(), before)
        response = self.client.post("/gestion-5s/import/ordenes/cancel", data=values)
        self.assertEqual(response.status_code, 302)
        self.assertIn("tab=miscelaneo", response.location)
        self.assertEqual(self.snapshot(), before)
        self.assertEqual(self.confirm(preview).status_code, 409)

    def test_confirmation_splits_exact_specialty_replaces_all_and_preserves_other_categories(self):
        before_others = {key: self.rows(key) for key in ("reclamos", "samtech_usuarios")}
        orders = [order("000001", "  carpintería  menor  "), order("000002", "CARPINTERIA"),
                  order("000003", "CARPINTERIA MENOR EXTRA"), order("000004", "")]
        response = self.confirm(self.preview(orders))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.startswith("/gestion-5s/registros?vista=solicitud_ot"))
        self.assertEqual([r["ot"] for r in self.rows("miscelaneo")], ["000001"])
        self.assertEqual([r["ot"] for r in self.rows("solicitud_ot")], ["000002", "000003", "000004"])
        self.assertEqual(self.rows("miscelaneo")[0]["especialidad"], "carpintería  menor")
        for key, rows in before_others.items():
            self.assertEqual(self.rows(key), rows)

    def test_same_file_can_replace_again_but_same_confirmation_cannot_repeat(self):
        source = [order(), order("000002", "CARPINTERIA MENOR")]
        preview = self.preview(source)
        self.assertEqual(self.confirm(preview).status_code, 302)
        first = self.snapshot()
        self.assertEqual(self.confirm(preview).status_code, 409)
        self.assertEqual(self.snapshot(), first)
        self.assertEqual(self.confirm(self.preview(source)).status_code, 302)
        self.assertEqual([len(self.rows(key)) for key in ORDER_ENTITIES], [1, 1])

    def test_invalid_empty_and_old_layout_files_never_delete_records(self):
        before = self.snapshot()
        bad_header = [label for _, label in ORDER_IMPORT_FIELDS]
        bad_header[6] = "Otra columna"
        files = [workbook_bytes([]), b"not an xlsx", workbook_bytes([order()], bad_header),
                 workbook_bytes([order(fecha_creacion="31/02/2026")]),
                 workbook_bytes([order(), order("2", fecha_aprobacion="fecha inválida")]),
                 workbook_bytes([order(estado="x" * 101)]), workbook_bytes([order(comentario="=1+1")])]
        for content in files:
            with self.subTest(length=len(content)):
                self.preview([], content=content, expected=400)
                self.assertEqual(self.snapshot(), before)

    def test_blanks_reordered_headers_excel_dates_long_text_and_duplicate_tickets_are_preserved(self):
        rows = [order("000001", "CARPINTERIA MENOR", fecha_creacion=None, fecha_inicio=None),
                order("000001", "ELECTRICO", fecha_creacion=None, fecha_inicio="12/09/2026"),
                order("000003", "GASFITER", fecha_creacion=date(2026, 8, 31), falla="f" * 600,
                      comentario="Texto\n" + "c" * 1000),
                [None] * 15]
        preview = self.preview([], content=workbook_bytes(rows, reverse=True))
        self.assertIn("1 repeticiones de ticket", preview.get_data(as_text=True))
        self.assertEqual(self.confirm(preview).status_code, 302)
        self.assertIsNone(self.rows("miscelaneo")[0]["fecha_creacion"])
        imported = self.rows("solicitud_ot")
        self.assertEqual(imported[0]["ot"], "000001")
        self.assertEqual(imported[0]["fecha_inicio"], date(2026, 9, 12))
        self.assertEqual(imported[1]["falla"], "f" * 600)
        self.assertEqual(imported[1]["comentario"], rows[2]["comentario"])
        numeric = workbook_bytes([order(fecha_creacion=to_excel(date(2026, 9, 12), CALENDAR_MAC_1904))], epoch=CALENDAR_MAC_1904)
        self.assertEqual(parse_orders(numeric)["groups"]["solicitud_ot"][0]["fecha_creacion"], date(2026, 9, 12))

    def test_category_with_zero_rows_is_cleared_only_after_explicit_confirmation(self):
        preview = self.preview([order()])
        self.assertIn("quedará vacía", preview.get_data(as_text=True))
        self.assertEqual(len(self.rows("miscelaneo")), 1)
        self.assertEqual(self.confirm(preview).status_code, 302)
        self.assertEqual(self.rows("miscelaneo"), [])

    def test_paginated_orders_keep_search_export_and_delete_all_on_the_full_result(self):
        orders = [order(f"PAGE-{i:04d}") for i in range(105)]
        self.assertEqual(self.confirm(self.preview(orders)).status_code, 302)
        query = "vista=solicitud_ot&q=PAGE&from=2026-08-31&to=2026-08-31"
        first = self.client.get("/gestion-5s/registros?" + query).get_data(as_text=True)
        second = self.client.get("/gestion-5s/registros?" + query + "&page=2").get_data(as_text=True)
        self.assertEqual(first.count('class="form-check-input record-select"'), 100)
        self.assertEqual(second.count('class="form-check-input record-select"'), 5)
        self.assertIn("Mostrando 101–105 de 105", second)
        self.assertIn("Eliminar todos (105)", second)
        self.assertIn("Seleccionar todos los registros de esta página", second)
        links = PageStructure(first).links
        next_page = next(link for link in links if "page=2" in link and "/registros?" in link)
        self.assertIn("q=PAGE", next_page)
        exported = self.client.get("/gestion-5s/download/solicitud_ot.csv?" + query + "&page=2")
        self.assertEqual(len(list(csv.DictReader(StringIO(exported.data.decode("utf-8-sig"))))), 105)
        with web.app.test_request_context(environ_base={"SCRIPT_NAME": "/gestion-5s"}):
            returned = web.records_return_url("solicitud_ot", "/registros?" + query + "&page=2")
            self.assertIn("page=2", returned)
        fields = FormValues(second).values
        confirm = self.client.post("/gestion-5s/delete/solicitud_ot/bulk/confirm", data={
            "csrf_token": fields["csrf_token"], "mode": "all", "q": "PAGE",
            "from": "2026-08-31", "to": "2026-08-31",
        })
        self.assertEqual(confirm.status_code, 200)
        payload = FormValues(confirm.get_data(as_text=True)).values
        result = self.client.post("/gestion-5s/delete/solicitud_ot/bulk", data=payload)
        self.assertEqual(result.status_code, 302)
        self.assertEqual(self.rows("solicitud_ot"), [])

    def test_concurrent_insert_edit_or_delete_requires_a_new_preview(self):
        for operation in ("insert", "edit", "delete"):
            with self.subTest(operation=operation):
                preview = self.preview([order()])
                with self.sessions() as db:
                    if operation == "insert":
                        db.add(web.MiscelaneoEntry(ot="concurrent"))
                    elif operation == "edit":
                        db.query(web.MiscelaneoEntry).first().comentario = "Cambio posterior"
                    else:
                        db.delete(db.query(web.MiscelaneoEntry).first())
                    db.commit()
                changed = self.snapshot()
                self.assertEqual(self.confirm(preview).status_code, 409)
                self.assertEqual(self.snapshot(), changed)

    def test_failure_after_delete_and_partial_insert_rolls_back_both_tables(self):
        before = self.snapshot()
        preview = self.preview([order(), order("2", "CARPINTERIA MENOR")])
        def fail_midway(db, models, groups):
            db.execute(models["miscelaneo"].__table__.insert(), groups["miscelaneo"])
            raise SQLAlchemyError("simulated failure")
        with patch("gestion5s.order_import_routes._insert_groups", side_effect=fail_midway), self.assertLogs(web.app.logger, level="ERROR"):
            self.assertEqual(self.confirm(preview).status_code, 503)
        self.assertEqual(self.snapshot(), before)

    def test_session_tampering_expiry_and_get_cannot_confirm(self):
        before = self.snapshot()
        preview = self.preview([order()])
        values = FormValues(preview.get_data(as_text=True)).values
        bad = {**values, "confirm_replace": "1", "confirmation_token": values["confirmation_token"] + "invalid"}
        self.assertEqual(self.client.post("/gestion-5s/import/ordenes/confirm", data=bad).status_code, 400)
        other = app.test_client()
        stolen = {**values, "confirm_replace": "1", "csrf_token": self.csrf(other)}
        self.assertEqual(other.post("/gestion-5s/import/ordenes/confirm", data=stolen).status_code, 400)
        self.assertEqual(self.client.post("/gestion-5s/import/solicitud_ot", data={"file": (BytesIO(workbook_bytes([order()])), "a.xlsx")}).status_code, 400)
        self.assertEqual(self.client.get("/gestion-5s/import/ordenes/confirm").status_code, 405)
        with self.sessions() as db:
            db.query(web.OrderImportBatch).first().expires_at = web.now_utc() - timedelta(seconds=1)
            db.commit()
        self.assertEqual(self.confirm(preview).status_code, 409)
        self.assertEqual(self.snapshot(), before)

    def test_report_uses_imported_creation_status_and_fallback_without_miscellaneous(self):
        orders = [order("M", "CARPINTERIA MENOR", estado="Aprobada"),
                  order("A", estado="Aprobada", fecha_inicio="15/10/2026"),
                  order("B", estado="Completada", fecha_creacion="12/09/2026"),
                  order("C", estado="No Iniciada", fecha_creacion="01/09/2026", fecha_inicio=None),
                  order("D", estado="En Progreso", fecha_creacion=None, fecha_inicio="12/09/2026"),
                  order("E", estado="Eliminado"), order("F", estado="Felicitaciones"),
                  order("G", fecha_creacion=None, fecha_inicio=None),
                  order("X", fecha_creacion="13/09/2026", fecha_inicio="01/09/2026")]
        self.assertEqual(self.confirm(self.preview(orders)).status_code, 302)
        report = self.report()
        category = next(c for c in report["categories"] if c["entity"] == "solicitud_ot")
        self.assertEqual({key: category["totals"][key] for key in ("records", "open", "closed", "unclassified")},
                         {"records": 6, "open": 2, "closed": 2, "unclassified": 2})
        self.assertEqual(category["counts"]["records"], [3, 1] + [0] * 10 + [2])
        self.assertEqual(category["undated"], 1)
        self.assertEqual(classify_status("Aprobada", "solicitud_ot"), "cerrado")
        self.assertEqual(classify_status("Aprobada", "samtech_usuarios"), "sin_clasificar")
        self.assertEqual(self.report(details=True)["counts"], report["counts"])
        filters = {"start_date": "2026-08-31", "end_date": "2026-09-12"}
        page = self.client.get("/reports/usuarios", query_string=filters).get_data(as_text=True)
        self.assertIn("Aprobada y Completada", page)
        response = self.client.get("/reports/usuarios.xlsx", query_string=filters)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), data_only=True)
        sheet = workbook["Solicitudes OT"]
        headers = [c.value for c in sheet[5]]
        rows = list(sheet.iter_rows(min_row=6, values_only=True))
        self.assertEqual(len(rows), 6)
        self.assertEqual({row[0] for row in rows}, {"A", "B", "C", "D", "E", "F"})
        self.assertEqual(headers[:15], [label for _, label in ORDER_IMPORT_FIELDS])
        self.assertEqual(headers[-2:], ["Fecha para reporte", "Estado agrupado"])
        self.assertEqual(sum(row[-1] == "Cerrado" for row in rows), 2)
        totals = next(row for row in workbook["Reporte diario"] if row[0].value == "Solicitudes de usuarios")
        self.assertEqual(totals[-1].value, 6)
        formula_book = load_workbook(BytesIO(response.data))
        self.assertTrue(formula_book["Solicitudes OT"].cell(6, len(headers) - 1).value.startswith("=IF("))
        # Listado, búsqueda, CSV y dashboard usan la misma fecha que Gestión de usuarios.
        url = "/gestion-5s/registros?vista=solicitud_ot&from=2026-08-31&to=2026-09-12"
        listing = self.client.get(url).get_data(as_text=True)
        self.assertIn("Fecha creación (Fecha inicio si falta)", listing)
        download = self.client.get("/gestion-5s/download/solicitud_ot.csv?from=2026-08-31&to=2026-09-12")
        exported = list(csv.DictReader(StringIO(download.data.decode("utf-8-sig"))))
        self.assertEqual({row["ot"] for row in exported}, {"A", "B", "C", "D", "E", "F"})
        self.assertIn("especialidad", exported[0])
        self.assertIn("comentario", exported[0])
        self.assertEqual(self.client.get("/gestion-5s/dashboard?from=2026-08-31&to=2026-09-12").status_code, 200)

    def test_migration_adds_fields_without_deleting_legacy_data_and_can_repeat(self):
        engine = create_engine("sqlite://")
        metadata = MetaData()
        legacy = {"id", "ot", "fecha_inicio", "estado", "creado", "n_solicitud", "observacion"}
        columns = [Column(c.name, c.type, primary_key=c.primary_key, nullable=c.nullable)
                   for c in web.SolicitudOTEntry.__table__.columns if c.name in legacy]
        Table("solicitudes_ot", metadata, *columns)
        web.MiscelaneoEntry.__table__.to_metadata(metadata)
        metadata.create_all(engine)
        with engine.begin() as conn:
            conn.execute(text("INSERT INTO solicitudes_ot (id,ot,fecha_inicio,estado,creado,observacion) VALUES (1,'LEGACY','2026-09-01','Abierto','2026-09-01','Conservar')"))
        web.ensure_order_import_columns(engine)
        web.ensure_order_import_columns(engine)
        self.assertTrue({name for name, _ in ORDER_IMPORT_FIELDS} <= {c["name"] for c in inspect(engine).get_columns("solicitudes_ot")})
        with engine.connect() as conn:
            row = conn.execute(text("SELECT ot,observacion,fecha_creacion FROM solicitudes_ot")).one()
            self.assertEqual(tuple(row), ("LEGACY", "Conservar", None))
        engine.dispose()


if __name__ == "__main__":
    unittest.main()
