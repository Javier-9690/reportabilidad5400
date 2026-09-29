from datetime import date, time, timedelta
from io import BytesIO
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from tests.csv_helpers import excel_csv_reader
from tests.test_integration import app
from tests.test_edit_records import FormValues
from tests.test_form_structure import PageStructure
from gestion5s import web
from gestion5s.alarm_times import ALARM_HOUR_FIELDS, ALARM_TIME_FIELDS


class AlarmTimesTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        web.Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        self.session_patch = patch.object(web, "SessionLocal", self.sessions)
        self.session_patch.start()
        self.client = app.test_client()

    def tearDown(self):
        self.session_patch.stop()
        self.engine.dispose()

    def records(self):
        with self.sessions() as db:
            return [{column.name: getattr(record, column.name) for column in record.__table__.columns}
                    for record in db.query(web.ActivacionAlarmaEntry).order_by(web.ActivacionAlarmaEntry.id)]

    def create(self, **values):
        response = self.client.post("/gestion-5s/panel?tab=alarmas", data={"fecha": "2026-09-27", **values})
        self.assertEqual(response.status_code, 302)
        return self.records()[-1]

    def editor(self, record):
        return self.client.get(f"/gestion-5s/edit/alarmas/{record['id']}").get_data(as_text=True)

    def exported(self, **filters):
        response = self.client.get("/gestion-5s/download/alarmas.csv", query_string=filters)
        self.assertEqual(response.status_code, 200)
        return list(excel_csv_reader(response.data))

    def upload(self, workbook):
        output = BytesIO()
        workbook.save(output)
        response = self.client.post("/gestion-5s/import/alarmas", data={"file": (BytesIO(output.getvalue()), "alarmas.xlsx")}, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        return response.get_data(as_text=True)

    def workbook(self, rows):
        book = Workbook()
        book.active.append(web.TEMPLATES["alarmas"])
        for values in rows:
            row = {"FECHA": "2026-09-27", **values}
            book.active.append([row.get(header) for header in web.TEMPLATES["alarmas"]])
        return book

    def test_all_five_fields_keep_seconds_through_create_list_export_and_edit(self):
        data = dict(zip(ALARM_TIME_FIELDS, ("08:15:43", "09:02:17", "09:03:59", "09:30:01", "09:45:08")))
        record = self.create(**data)
        for page in (self.client.get("/gestion-5s/panel?tab=alarmas").get_data(as_text=True), self.editor(record)):
            controls = {control['name']: control for form in PageStructure(page).forms for control in form['controls']}
            for field in ALARM_TIME_FIELDS:
                self.assertNotIn("required", controls[field])
                if field in ALARM_HOUR_FIELDS:
                    self.assertEqual(controls[field]["type"], "text")
                    self.assertEqual(controls[field]["placeholder"], "08:15:30")
                    self.assertIn("[0-5][0-9]", controls[field]["pattern"])
                else:
                    self.assertEqual((controls[field]["type"], controls[field]["step"]), ("time", "1"))
        editor = FormValues(self.editor(record)).values
        listing = self.client.get("/gestion-5s/registros?vista=alarmas").get_data(as_text=True)
        cells = PageStructure(listing).cells
        exported = self.exported()[0]
        for field, value in data.items():
            self.assertEqual(editor[field], value)
            self.assertEqual(cells[field], value)
            self.assertEqual(exported[field.upper()], value)
        self.assertEqual(record["hora_reporte_salfa"], time(9, 45, 8))
        for field in ALARM_HOUR_FIELDS:
            h, m, s = map(int, data[field].split(":"))
            self.assertAlmostEqual(record[field], (h * 3600 + m * 60 + s) / 3600)
        edited = {**editor, **dict.fromkeys(ALARM_TIME_FIELDS, "11:22:33")}
        response = self.client.post(f"/gestion-5s/edit/alarmas/{record['id']}", data=edited)
        self.assertEqual(response.status_code, 302)
        self.assertEqual((self.records()[0]["id"], self.records()[0]["creado"]), (record["id"], record["creado"]))
        for field in ALARM_TIME_FIELDS:
            self.assertEqual(self.exported()[0][field.upper()], "11:22:33")

    def test_empty_values_and_midnight_remain_distinct_and_can_be_cleared(self):
        empty = self.create()
        self.assertTrue(all(empty[field] is None for field in ALARM_TIME_FIELDS))
        midnight = self.create(**dict.fromkeys(ALARM_TIME_FIELDS, "00:00:00"))
        self.assertTrue(all(midnight[field] == 0 for field in ALARM_HOUR_FIELDS))
        self.assertEqual(midnight["hora_reporte_salfa"], time(0, 0))
        exported = self.exported()
        self.assertTrue(all(exported[0][field.upper()] == "" for field in ALARM_TIME_FIELDS))
        self.assertTrue(all(exported[1][field.upper()] == "00:00:00" for field in ALARM_TIME_FIELDS))
        form = FormValues(self.editor(midnight)).values
        form.update(dict.fromkeys(ALARM_TIME_FIELDS, ""))
        self.assertEqual(self.client.post(f"/gestion-5s/edit/alarmas/{midnight['id']}", data=form).status_code, 302)
        self.assertTrue(all(self.records()[1][field] is None for field in ALARM_TIME_FIELDS))

    def test_historical_decimal_hours_and_precision_are_preserved_when_editing_another_field(self):
        values = {"aviso_mantencion_h": 0.5, "llegada_mantencion_h": 1.25,
                  "aviso_lider_h": 1.234567, "llegada_lider_h": 27.5125}
        with self.sessions() as db:
            db.add(web.ActivacionAlarmaEntry(fecha=date(2026, 9, 27), hora_reporte_salfa=time(23, 59, 57), **values))
            db.commit()
        record = self.records()[0]
        form = FormValues(self.editor(record)).values
        self.assertEqual([form[field] for field in ALARM_HOUR_FIELDS], ["00:30:00", "01:15:00", "01:14:04", "27:30:45"])
        form["observaciones"] = "Revisión sin cambiar horarios"
        self.assertEqual(self.client.post(f"/gestion-5s/edit/alarmas/{record['id']}", data=form).status_code, 302)
        for field, value in values.items():
            self.assertEqual(self.records()[0][field], value)
        self.assertEqual(self.exported()[0]["HORA_REPORTE_SALFA"], "23:59:57")

    def test_template_has_time_formats_and_no_example_that_could_be_imported(self):
        response = self.client.get("/gestion-5s/template/alarmas.xlsx")
        workbook = load_workbook(BytesIO(response.data))
        sheet = workbook.active
        self.assertEqual([c.value for c in sheet[1]], web.TEMPLATES["alarmas"])
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertFalse(any(value is not None for row in sheet.iter_rows(min_row=2, values_only=True) for value in row))
        for column in range(8, 13):
            self.assertIn("mm:ss", sheet.cell(2, column).number_format)
            self.assertIn("08:15:30", sheet.cell(1, column).comment.text)
        before = self.records()
        self.assertIn("OK: 0 filas", self.upload(workbook))
        self.assertEqual(self.records(), before)
        for column, header in enumerate(web.TEMPLATES["alarmas"], 1):
            if header == "FECHA":
                sheet.cell(2, column, date(2026, 9, 27))
            elif 8 <= column <= 12:
                sheet.cell(2, column, time(18, 42, 31))
        self.assertIn("OK: 1 filas", self.upload(workbook))
        self.assertTrue(all(self.exported()[0][field.upper()] == "18:42:31" for field in ALARM_TIME_FIELDS))

    def test_import_accepts_text_excel_times_blanks_and_legacy_decimal_hours(self):
        book = self.workbook([
            {**dict.fromkeys((f.upper() for f in ALARM_HOUR_FIELDS), "23:59:59"), "HORA_REPORTE_SALFA": "23:59:59"},
            {"AVISO_MANTENCION_H": time(8, 15, 43), "LLEGADA_MANTENCION_H": timedelta(hours=27, minutes=30, seconds=45),
             "AVISO_LIDER_H": 0.5, "LLEGADA_LIDER_H": "1,25", "HORA_REPORTE_SALFA": "10:40:21 PM"},
            {},
            {**dict.fromkeys((f.upper() for f in ALARM_TIME_FIELDS), 0)},
            {"AVISO_MANTENCION_H": "08:15", "HORA_REPORTE_SALFA": "09:30"},
        ])
        self.assertIn("OK: 5 filas", self.upload(book))
        rows = self.exported()
        self.assertTrue(all(rows[0][field.upper()] == "23:59:59" for field in ALARM_TIME_FIELDS))
        self.assertEqual([rows[1][field.upper()] for field in ALARM_TIME_FIELDS], ["08:15:43", "27:30:45", "00:30:00", "01:15:00", "22:40:21"])
        self.assertTrue(all(rows[2][field.upper()] == "" for field in ALARM_TIME_FIELDS))
        self.assertTrue(all(rows[3][field.upper()] == "00:00:00" for field in ALARM_TIME_FIELDS))
        self.assertEqual((rows[4]["AVISO_MANTENCION_H"], rows[4]["HORA_REPORTE_SALFA"]), ("08:15:00", "09:30:00"))

    def test_invalid_times_reject_form_and_roll_back_the_whole_import(self):
        record = self.create(aviso_mantencion_h="08:15:30")
        before = self.records()
        for field, value in (("aviso_mantencion_h", "08:60:00"), ("llegada_mantencion_h", "08:15:60"),
                             ("aviso_lider_h", "NaN"), ("llegada_lider_h", "-1"),
                             ("hora_reporte_salfa", "25:00:00"), ("hora_reporte_salfa", "mal escrito")):
            with self.subTest(field=field, value=value):
                form = FormValues(self.editor(record)).values
                response = self.client.post(f"/gestion-5s/edit/alarmas/{record['id']}", data={**form, field: value})
                self.assertEqual(response.status_code, 422)
                self.assertEqual(self.records(), before)
                self.assertEqual(self.client.post("/gestion-5s/panel?tab=alarmas", data={"fecha": "2026-09-27", field: value}).status_code, 422)
                page = self.upload(self.workbook([{"AVISO_MANTENCION_H": "11:22:33"}, {field.upper(): value}]))
                self.assertIn("Fila 3", page)
                self.assertIn("Error importando alarmas", page)
                self.assertEqual(self.records(), before)

    def test_hms_search_matches_list_export_and_filtered_delete_without_touching_other_rows(self):
        selected = self.create(llegada_lider_h="19:43:27")
        other = self.create(llegada_lider_h="19:43:28")
        page = self.client.get("/gestion-5s/registros", query_string={"vista": "alarmas", "q": "19:43:27"}).get_data(as_text=True)
        self.assertIn(f"/edit/alarmas/{selected['id']}?", page)
        self.assertNotIn(f"/edit/alarmas/{other['id']}?", page)
        self.assertEqual(len(self.exported(q="19:43:27")), 1)
        self.assertEqual(len(self.exported(q=":43:27")), 1)
        preview = self.client.post("/gestion-5s/delete/alarmas/bulk/confirm", data={
            "csrf_token": FormValues(page).values["csrf_token"], "mode": "all", "q": "19:43:27",
        })
        self.assertEqual(preview.status_code, 200)
        result = self.client.post("/gestion-5s/delete/alarmas/bulk", data=FormValues(preview.get_data(as_text=True)).values)
        self.assertEqual(result.status_code, 302)
        self.assertEqual([row['id'] for row in self.records()], [other['id']])


if __name__ == "__main__":
    unittest.main()
