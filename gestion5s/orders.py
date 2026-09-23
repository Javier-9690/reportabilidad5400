"""Lectura completa y clasificación de órdenes antes de reemplazar registros."""

from collections import Counter
from datetime import date, datetime
from io import BytesIO
import math
import re
import unicodedata
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import func

from gestion5s.editing import ORDER_IMPORT_FIELDS

ORDER_ENTITIES = ("miscelaneo", "solicitud_ot")
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ORDER_ROWS = 150000
ORDER_REPORT_NOTE = (
    "Solicitudes de usuarios se alimenta de Solicitudes OT, sin CARPINTERIA MENOR, "
    "que se guarda en Misceláneos. Se usa Fecha creación y, si falta, Fecha inicio. "
    "Aprobada y Completada cuentan como cerradas en Solicitudes OT; Eliminado y "
    "Felicitaciones quedan Sin clasificar. Samtech usuarios conserva su registro independiente."
)


def normalize_order_text(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).upper().split())


def order_header(value):
    text = normalize_order_text(value).replace("_", " ")
    return "TICKET" if text == "OT" else " ".join(text.split())


def order_destination(specialty):
    return "miscelaneo" if normalize_order_text(specialty) == "CARPINTERIA MENOR" else "solicitud_ot"


def order_reference_date(record):
    return record.fecha_creacion or record.fecha_inicio


def order_reference_column(model):
    return func.coalesce(model.fecha_creacion, model.fecha_inicio)


def _blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def _date(value, epoch):
    if _blank(value):
        return None
    if isinstance(value, bool):
        raise ValueError("fecha inválida")
    if isinstance(value, (int, float)):
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        if value < date(1900, 1, 1):
            raise ValueError("fecha anterior a 1900")
        return value
    text = str(value).strip()
    for pattern in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S",
                    "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return _date(datetime.strptime(text, pattern), epoch)
        except ValueError:
            pass
    raise ValueError("usa una fecha válida (dd/mm/aaaa o aaaa-mm-dd), o deja la celda vacía")


def _text(cell, identifier=False):
    value = cell.value
    if _blank(value):
        return None
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("número inválido")
        value = int(value) if value.is_integer() else value
    if identifier and isinstance(value, int) and re.fullmatch(r"0+", cell.number_format or ""):
        return str(value).zfill(len(cell.number_format))
    return str(value).strip()


def parse_orders(content):
    """No escribe en BD. Si una fila es inválida, se rechaza el archivo completo."""
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("Sube un archivo .xlsx de hasta 25 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("El contenido del Excel supera 200 MB. Divide el archivo antes de importar.")
    except BadZipFile:
        raise ValueError("El archivo no es un Excel .xlsx válido.") from None
    book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        expected = {order_header(label): name for name, label in ORDER_IMPORT_FIELDS}
        candidates = []
        for sheet in book:
            headers = [order_header(cell.value) for cell in next(sheet.iter_rows(max_row=1), ())]
            nonempty = [header for header in headers if header]
            if len(nonempty) == len(expected) and set(nonempty) == set(expected):
                candidates.append((sheet, {expected[header]: index for index, header in enumerate(headers) if header}))
        if len(candidates) != 1:
            if len(candidates) > 1:
                raise ValueError("Hay varias hojas de órdenes. Deja una sola hoja de datos para evitar omisiones.")
            raise ValueError("Se necesita una hoja con estos 15 encabezados, sin columnas repetidas: " +
                             ", ".join(label for _, label in ORDER_IMPORT_FIELDS) + ".")
        sheet, columns = candidates[0]
        groups = {entity: [] for entity in ORDER_ENTITIES}
        states = {entity: Counter() for entity in ORDER_ENTITIES}
        undated = dict.fromkeys(ORDER_ENTITIES, 0)
        fallback = dict.fromkeys(ORDER_ENTITIES, 0)
        dates, tickets = [], Counter()
        blank_rows = blank_specialty = total = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=2), 2):
            if all(_blank(cell.value) for cell in row):
                blank_rows += 1
                continue
            total += 1
            if total > MAX_ORDER_ROWS:
                raise ValueError(f"El archivo supera {MAX_ORDER_ROWS:,} órdenes. No se reemplazó ningún registro.")
            values = {}
            for name, label in ORDER_IMPORT_FIELDS:
                cell = row[columns[name]]
                try:
                    if cell.data_type in ("f", "e"):
                        raise ValueError("reemplaza la fórmula o el error por un valor")
                    if name.startswith("fecha_"):
                        value = _date(cell.value, book.epoch)
                    else:
                        value = _text(cell, identifier=name == "ot")
                        limit = 100 if name in ("ot", "estado") else 200
                        if name not in ("falla", "comentario") and value and len(value) > limit:
                            raise ValueError(f"máximo {limit} caracteres")
                    values[name] = value
                except (ValueError, TypeError, OverflowError) as exc:
                    raise ValueError(f"Fila {row_number}, {label}: {exc}.") from exc
            if any(not _blank(cell.value) for i, cell in enumerate(row) if i not in columns.values()):
                raise ValueError(f"Fila {row_number}: hay datos en una columna sin encabezado.")
            entity = order_destination(values["especialidad"])
            groups[entity].append(values)
            states[entity][values["estado"] or "(vacío)"] += 1
            if not values["especialidad"]:
                blank_specialty += 1
            day = values["fecha_creacion"] or values["fecha_inicio"]
            if day:
                dates.append(day)
                if not values["fecha_creacion"]:
                    fallback[entity] += 1
            else:
                undated[entity] += 1
            if values["ot"]:
                tickets[values["ot"]] += 1
        if not total:
            raise ValueError("El archivo no contiene órdenes. No se permite vaciar ambas bases con una plantilla vacía.")
        return {"groups": groups, "sheet": sheet.title, "total": total, "blank_rows": blank_rows,
                "states": states, "undated": undated, "fallback": fallback, "blank_specialty": blank_specialty,
                "duplicate_tickets": sum(count - 1 for count in tickets.values()),
                "start": min(dates) if dates else None, "end": max(dates) if dates else None}
    finally:
        book.close()
