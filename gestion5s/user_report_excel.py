"""Exportación del informe desde Flask con las dependencias del programa."""

from datetime import date, datetime, time
from io import BytesIO
import math

import xlsxwriter
from xlsxwriter.utility import xl_col_to_name, xl_rowcol_to_cell
from openpyxl.utils.datetime import to_excel

from gestion5s.user_reports import OPEN_STATUSES, STATUS_ALIASES, STATUS_LABELS, classify_status


def count_axis_step(maximum):
    target = max(1, math.ceil(maximum / 6))
    base = 10 ** int(math.log10(target))
    return next(multiplier * base for multiplier in (1, 2, 5, 10) if multiplier * base >= target)


def export_user_report(report):
    for category in report["categories"]:
        if len(category["details"]) + 5 > 1048576:
            raise ValueError("El detalle supera las filas disponibles en Excel. Selecciona un período más corto.")
        for record in category["details"]:
            if any(isinstance(value, str) and len(value) > 32767 for value in record.values()):
                raise ValueError(f"Un texto de {category['label']} supera los 32.767 caracteres admitidos por celda de Excel.")

    output = BytesIO()
    book = xlsxwriter.Workbook(output, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False})
    book.set_properties({"title": "Gestión de usuarios — Reportabilidad 5400", "author": "Reportabilidad 5400"})
    book.set_calc_mode("auto")
    base = {"font_name": "Calibri", "font_size": 11, "valign": "vcenter"}

    def fmt(**properties):
        return book.add_format({**base, **properties})

    formats = {
        "title": fmt(font_size=18, bold=True, font_color="#FFFFFF", bg_color="#B42318", text_wrap=True),
        "subtitle": fmt(font_color="#555555", text_wrap=True),
        "header": fmt(bold=True, font_color="#FFFFFF", bg_color="#F00000", border=1, text_wrap=True),
        "date_header": fmt(bold=True, font_color="#FFFFFF", bg_color="#F00000", border=1, num_format="[$-340A]d-mmm", align="right"),
        "section": fmt(bold=True, font_color="#20473B", font_size=12),
        "label": fmt(bold=True, border=1),
        "number": fmt(border=1, num_format="#,##0", align="right"),
        "total_label": fmt(bold=True, font_color="#FFFFFF", bg_color="#F00000", border=1),
        "total": fmt(bold=True, bg_color="#FFFF00", border=1, num_format="#,##0", align="right"),
        "period": fmt(bold=True, bg_color="#FFF3C2", border=1, num_format="#,##0", align="right"),
        "percent": fmt(bold=True, border=1, num_format="0%", align="right"),
        "period_percent": fmt(bold=True, bg_color="#FFF3C2", border=1, num_format="0%", align="right"),
        "text": fmt(text_wrap=True),
        "source_text": fmt(text_wrap=True, num_format="@"),
        "date": fmt(num_format="dd/mm/yyyy"),
        "datetime": fmt(num_format="dd/mm/yyyy hh:mm:ss"),
        "time": fmt(num_format="hh:mm:ss"),
        "duration": fmt(num_format="[mm]:ss"),
        "raw_number": fmt(num_format="0.##"),
        "card_label": fmt(bold=True, font_color="#555555", bg_color="#F5F6F8", align="center", text_wrap=True),
        "card_value": fmt(bold=True, font_size=25, bg_color="#F5F6F8", align="center", num_format="#,##0"),
        "card_rate": fmt(bold=True, font_size=25, bg_color="#F5F6F8", align="center", num_format="0.0%"),
        "summary_rate": fmt(border=1, num_format="0.0%", align="right"),
    }
    daily = book.add_worksheet("Reporte diario")
    dashboard = book.add_worksheet("Conclusiones")
    source_ranges = {}
    period = f"Período: {report['start']:%d/%m/%Y} al {report['end']:%d/%m/%Y}"
    stamp = f"Datos al exportar: {report['generated_at']:%d/%m/%Y %H:%M} UTC"

    for category in report["categories"]:
        sheet = book.add_worksheet(category["sheet"])
        fields = category["fields"]
        last_column = len(fields)
        sheet.hide_gridlines(2)
        sheet.set_tab_color("#BFC5CA")
        sheet.merge_range(0, 0, 0, last_column, category["label"], formats["title"])
        sheet.set_row(0, 32)
        sheet.merge_range(1, 0, 1, last_column, period, formats["subtitle"])
        sheet.merge_range(2, 0, 2, last_column,
                          f"Fuente: Registros hotelería > {category['label']}. Fecha utilizada: {category['date_label']}. {stamp}", formats["subtitle"])
        sheet.set_row(2, 30)
        sheet.merge_range(3, 0, 3, last_column,
                          "Estado agrupado aplica las equivalencias del informe. Para incorporar cambios del programa, genera una nueva exportación.", formats["subtitle"])
        sheet.set_row(3, 30)
        for column, field in enumerate(fields):
            sheet.write_string(4, column, field["label"], formats["header"])
            sheet.set_column(column, column, 48 if field["kind"] == "textarea" else 23)
        sheet.write_string(4, last_column, "Estado agrupado", formats["header"])
        sheet.set_column(last_column, last_column, 22)
        sheet.set_row(4, 32)
        for row, record in enumerate(category["details"], 5):
            for column, field in enumerate(fields):
                value = record[field["name"]]
                if field["name"] == category.get("report_date_field"):
                    creation = xl_rowcol_to_cell(row, next(i for i, f in enumerate(fields) if f["name"] == category["date"]))
                    started = xl_rowcol_to_cell(row, next(i for i, f in enumerate(fields) if f["name"] == category["date_fallback"]))
                    sheet.write_formula(row, column, f'=IF({creation}<>"",{creation},IF({started}<>"",{started},""))',
                                        formats["date"], to_excel(value) if value else "")
                elif value is None:
                    sheet.write_blank(row, column, None, formats["text"])
                elif field["kind"] == "duration":
                    sheet.write_number(row, column, value / 86400, formats["duration"])
                elif isinstance(value, datetime):
                    sheet.write_datetime(row, column, value, formats["datetime"])
                elif isinstance(value, date):
                    sheet.write_datetime(row, column, value, formats["date"])
                elif isinstance(value, time):
                    sheet.write_number(row, column, (value.hour * 3600 + value.minute * 60 + value.second) / 86400, formats["time"])
                elif isinstance(value, (int, float)):
                    sheet.write_number(row, column, value, formats["raw_number"])
                else:
                    sheet.write_string(row, column, str(value), formats["source_text"])
            status_label = (STATUS_LABELS[classify_status(record[category["status"]], entity=category["entity"])]
                            if category["status"] else "Sin campo de estado")
            sheet.write_string(row, last_column, status_label, formats["text"])
        last_row = max(5, len(category["details"]) + 4)
        sheet.autofilter(4, 0, last_row, last_column)
        sheet.freeze_panes(5, 0)
        sheet.set_landscape()
        sheet.repeat_rows(4)
        sheet.set_footer("&LReportabilidad 5400&R&P / &N")
        date_index = next(index for index, field in enumerate(fields) if field["name"] == category.get("report_date_field", category["date"]))
        date_letter = xl_col_to_name(date_index)
        status_letter = xl_col_to_name(last_column)
        source_ranges[category["entity"]] = {
            "date": f"'{category['sheet']}'!${date_letter}$6:${date_letter}${last_row + 1}",
            "status": f"'{category['sheet']}'!${status_letter}$6:${status_letter}${last_row + 1}",
        }

    total_column = len(report["days"]) + 1
    title_end = min(total_column, 14)
    daily.hide_gridlines(2)
    daily.set_tab_color("#F00000")
    daily.set_column(0, 0, 51)
    daily.set_column(1, total_column - 1, 11)
    daily.set_column(total_column, total_column, 17)
    daily.merge_range(0, 0, 0, title_end, "GESTIÓN DE USUARIOS — CAMPAMENTO 5400", formats["title"])
    daily.set_row(0, 48 if total_column < 5 else 34)
    daily.merge_range(1, 0, 1, title_end, period + " · " + stamp, formats["subtitle"])
    daily.set_row(1, 32)
    daily.merge_range(2, 0, 2, title_end, report["method_note"], formats["subtitle"])
    daily.set_row(2, 45 if total_column < 5 else 30)
    daily.merge_range(3, 0, 3, title_end, report["rate_note"], formats["subtitle"])
    daily.set_row(3, 60 if total_column < 5 else 35)
    row_index = {}
    current_row = 5
    first_header = None

    for section in report["sections"]:
        daily.merge_range(current_row, 0, current_row, title_end, section["title"].upper(), formats["section"])
        daily.set_row(current_row, 25)
        current_row += 1
        header_row = current_row
        first_header = header_row if first_header is None else first_header
        daily.write_string(header_row, 0, "ÍTEM", formats["header"])
        for column, day in enumerate(report["days"], 1):
            daily.write_datetime(header_row, column, day, formats["date_header"])
        daily.write_string(header_row, total_column, "Total período", formats["header"])
        current_row += 1
        for item in section["rows"]:
            row_index[item["key"]] = current_row
            label_format = formats["total_label"] if item["kind"] == "total" else formats["label"]
            daily.write_string(current_row, 0, item["label"], label_format)
            daily.set_row(current_row, 20)
            for column, cached in enumerate(item["values"], 1):
                value_format = formats["percent"] if item["kind"] == "percent" else formats["total"] if item["kind"] == "total" else formats["number"]
                if item["kind"] == "percent":
                    opened = xl_rowcol_to_cell(row_index["total_open"], column)
                    closed = xl_rowcol_to_cell(row_index["total_closed"], column)
                    formula = f"=IF(SUM({opened},{closed})=0,0,{closed}/SUM({opened},{closed}))"
                elif item["refs"]:
                    references = [xl_rowcol_to_cell(row_index[key], column) for key in item["refs"]]
                    formula = "=" + references[0] if len(references) == 1 else "=SUM(" + ",".join(references) + ")"
                else:
                    category = next(c for c in report["categories"] if item["key"].startswith(c["entity"] + "_"))
                    kind = item["key"][len(category["entity"]) + 1:]
                    ranges = source_ranges[category["entity"]]
                    criteria = f"{ranges['date']},{xl_rowcol_to_cell(header_row, column, row_abs=True)}"
                    if kind != "records":
                        code = {"closed": "cerrado", "unclassified": "sin_clasificar"}.get(kind, kind)
                        status_label = STATUS_LABELS[code] if category["status"] else "Sin campo de estado"
                        criteria += f',{ranges["status"]},"{status_label}"'
                    formula = f"=COUNTIFS({criteria})"
                daily.write_formula(current_row, column, formula, value_format, cached)
            if item["kind"] == "percent":
                opened = xl_rowcol_to_cell(row_index["total_open"], total_column)
                closed = xl_rowcol_to_cell(row_index["total_closed"], total_column)
                formula = f"=IF(SUM({opened},{closed})=0,0,{closed}/SUM({opened},{closed}))"
                cell_format = formats["period_percent"]
            else:
                formula = f"=SUM(B{current_row + 1}:{xl_col_to_name(total_column - 1)}{current_row + 1})"
                cell_format = formats["total"] if item["kind"] == "total" else formats["period"]
            daily.write_formula(current_row, total_column, formula, cell_format, item["total"])
            current_row += 1
        current_row += 1
    daily.freeze_panes(first_header + 1, 1)
    daily.set_landscape()
    daily.set_paper(8)  # A3 permite leer el formato diario sin reducirlo a letra diminuta.
    if len(report["days"]) <= 3:
        daily.set_portrait()
        daily.set_paper(9)
        daily.fit_to_pages(1, 0)
    elif len(report["days"]) <= 14:
        daily.fit_to_pages(1, 1)
    else:
        daily.set_print_scale(90)
    daily.repeat_rows(0, first_header)
    daily.repeat_columns(0)
    daily.print_area(0, 0, current_row - 1, total_column)
    daily.set_footer("&LReportabilidad 5400&R&P / &N")

    def daily_ref(key, column=total_column):
        return "'Reporte diario'!" + xl_rowcol_to_cell(row_index[key], column, row_abs=True, col_abs=True)

    def merged_value(row, first, last, value, cell_format, cached=None, end_row=None):
        dashboard.merge_range(row, first, row if end_row is None else end_row, last, "", cell_format)
        if isinstance(value, str) and value.startswith("="):
            dashboard.write_formula(row, first, value, cell_format, cached)
        else:
            dashboard.write(row, first, value, cell_format)

    dashboard.hide_gridlines(2)
    dashboard.set_tab_color("#198754")
    dashboard.set_column(0, 18, 9)
    dashboard.merge_range("A1:S1", "CONCLUSIONES — GESTIÓN DE USUARIOS", formats["title"])
    dashboard.set_row(0, 34)
    dashboard.merge_range("A2:S2", period, formats["subtitle"])
    dashboard.merge_range("A3:S3", stamp, formats["subtitle"])
    cards = [("records", "Total registros", "total_records"), ("open", "Abiertos", "total_open"),
             ("closed", "Cerrados", "total_closed"), ("unclassified", "Sin clasificar", "total_unclassified"),
             ("rate", "% cerrado", "closure_rate")]
    for index, (key, label, source_key) in enumerate(cards):
        first, last = index * 4, index * 4 + 2
        merged_value(4, first, last, label, formats["card_label"])
        formula = "=" + daily_ref(source_key) if source_key in row_index else "=0"
        merged_value(5, first, last, formula, formats["card_rate"] if key == "rate" else formats["card_value"],
                     cached=report["totals"][key], end_row=6)
    dashboard.set_row(4, 28)
    dashboard.set_row(5, 26)
    dashboard.set_row(6, 20)
    dashboard.merge_range("A9:S9", "Resumen por categoría", formats["section"])
    summary_columns = [(0, 5, "Categoría"), (6, 7, "Registros"), (8, 9, "Abiertos"),
                       (10, 11, "Cerrados"), (12, 14, "Sin clasificar"), (15, 18, "% cerrado")]
    for first, last, label in summary_columns:
        merged_value(9, first, last, label, formats["header"])
    for row, category in enumerate(report["categories"], 10):
        merged_value(row, 0, 5, category["label"], formats["label"])
        for key, first, last in (("records", 6, 7), ("open", 8, 9), ("closed", 10, 11), ("unclassified", 12, 14)):
            if key == "open":
                refs = [daily_ref(category["entity"] + "_" + code) for code in OPEN_STATUSES
                        if category["entity"] + "_" + code in row_index]
                formula = "=SUM(" + ",".join(refs) + ")" if refs else "=0"
            elif category["entity"] + "_" + key in row_index:
                formula = "=" + daily_ref(category["entity"] + "_" + key)
            else:
                formula = "=0"
            merged_value(row, first, last, formula, formats["number"], cached=category["totals"][key])
        if category["status"]:
            merged_value(row, 15, 18, f"=IF(SUM(I{row+1},K{row+1})=0,0,K{row+1}/SUM(I{row+1},K{row+1}))",
                         formats["summary_rate"], cached=category["totals"]["rate"])
        else:
            merged_value(row, 15, 18, "No aplica", formats["summary_rate"])
        dashboard.set_row(row, 25)
    summary_last_row = 10 + len(report["categories"]) - 1
    conclusions_row = summary_last_row + 3
    dashboard.merge_range(conclusions_row, 0, conclusions_row, 18, "Conclusiones del período", formats["section"])
    for row, conclusion in enumerate(report["conclusions"], conclusions_row + 1):
        dashboard.merge_range(row, 0, row, 18, conclusion, formats["text"])
        dashboard.set_row(row, 32)
    chart_row = conclusions_row + 2 + len(report["conclusions"])
    trend = book.add_chart({"type": "line"})
    for category in report["categories"]:
        trend.add_series({"name": category["label"], "categories": ["Reporte diario", first_header, 1, first_header, total_column - 1],
                          "values": ["Reporte diario", row_index[category["entity"] + "_records"], 1,
                                     row_index[category["entity"] + "_records"], total_column - 1],
                          "line": {"color": category["color"], "width": 2}, "smooth": False})
    trend.set_title({"name": "Registros por día"})
    trend.set_x_axis({"date_axis": True, "num_format": "[$-340A]d-mmm"})
    daily_max = max(max(c["counts"]["records"]) for c in report["categories"])
    trend.set_y_axis({"min": 0, "num_format": "0", "major_unit": count_axis_step(daily_max)})
    trend.set_legend({"position": "bottom", "font": {"size": 9}})
    trend.set_size({"width": 580, "height": 320})
    dashboard.insert_chart(chart_row, 0, trend)
    status = book.add_chart({"type": "bar", "subtype": "stacked"})
    for label, column, color in (("Abiertos", 8, "#CA8100"), ("Cerrados", 10, "#198754"), ("Sin clasificar", 12, "#777F87")):
        status.add_series({"name": label, "categories": ["Conclusiones", 10, 0, summary_last_row, 0],
                           "values": ["Conclusiones", 10, column, summary_last_row, column],
                           "fill": {"color": color}, "border": {"none": True}})
    status.set_title({"name": "Estado actual por categoría"})
    status.set_x_axis({"min": 0, "num_format": "0", "major_unit": count_axis_step(max(c["totals"]["records"] for c in report["categories"]))})
    status.set_legend({"position": "bottom", "font": {"size": 9}})
    status.set_size({"width": 580, "height": 320})
    dashboard.insert_chart(chart_row, 10, status)
    note_row = chart_row + 18
    dashboard.merge_range(note_row, 0, note_row, 18, "Criterios del informe", formats["section"])
    notes = [report["method_note"], report["rate_note"], report["without_status_note"], report["orders_note"],
             "Cada fila guardada cuenta una vez. Los tickets repetidos cuentan por separado. Se incluyen todos los días del rango.",
             "Fechas utilizadas: " + "; ".join(f"{c['label']}: {c['date_label']}" for c in report["categories"]),
             "Registros sin fecha en el histórico, excluidos del rango: " + "; ".join(f"{c['label']}: {c['undated']}" for c in report["categories"])]
    notes += [f"{STATUS_LABELS[key]}: {', '.join(aliases)}." for key, aliases in STATUS_ALIASES.items()]
    notes.append("Se ignoran mayúsculas y tildes. Los estados vacíos u otros valores, incluidos cancelado y rechazado, quedan Sin clasificar.")
    notes += [f"Estado sin clasificar — {c['label']}: {item['state']} ({item['count']} registros del período)."
              for c in report["categories"] for item in c["unknown_states"]]
    for row, note in enumerate(notes, note_row + 1):
        dashboard.merge_range(row, 0, row, 18, note, formats["subtitle"])
        dashboard.set_row(row, 32)
    dashboard.set_landscape()
    dashboard.set_paper(8)
    dashboard.set_print_scale(85)
    # Los gráficos deben empezar completos en una página al imprimir el resumen
    # de las cinco categorías; las conclusiones pueden ocupar más de media hoja.
    dashboard.set_h_pagebreaks([chart_row])
    dashboard.repeat_rows(0, 2)
    dashboard.print_area(0, 0, note_row + len(notes), 18)
    dashboard.set_footer("&LReportabilidad 5400&R&P / &N")
    daily.activate()
    book.close()
    output.seek(0)
    return output
